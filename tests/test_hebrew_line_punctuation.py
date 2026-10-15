"""Session 381: a quoted Hebrew line keeps its own punctuation at its END.

The bug: `_segment_by_script` walked the text character by character and put every
non-Hebrew character in an LTR run. On a line that is nothing but quoted Hebrew, that
split the sentence's own final `.` / `,` / `;` / `?` off into its own LTR run, and
Word's bidi engine then resolved that mark at the paragraph's LTR level and drew it at
the RIGHT edge of the island — hard against the line's FIRST word, so it read as if it
had OPENED the sentence.

Measured in Word's own PDF output for Psalm 74, Amichai's `בְּעֶרֶב שַׁבָּת.`:

    before   Hebrew run x=108.05   period drawn separately at x=155.83   (after בְּעֶרֶב)
    after    Hebrew run x=108.05   period is the run's leading glyph     (after שַׁבָּת)

502 such lines across 38 finished guides. A matched A/B of the whole Psalm 74 document
(same inputs, generator patched to withhold the verdict) moved exactly the two pages
carrying quoted Hebrew blocks and left the page count and all prose byte-identical.

What is pinned here:

1. The verdict is a property of the whole LINE, never of a markdown fragment. Splitting
   `the lot (גּוֹרָל), and` on its emphasis markers yields the Latin-free fragment
   ` (גּוֹרָל), `; treating THAT as a quoted Hebrew line drags the English sentence's
   comma inside the RTL run and renders it before the opening parenthesis. This is not
   hypothetical — it is what the first draft of the fix did, caught by diffing the
   fragments a real 75-guide build actually passes to the segmenter.

2. MIXED prose is untouched. There the mark after a Hebrew term belongs to the English
   sentence (`אֱלֹקִים, "God"`), and the island's right edge is where an English reader
   expects it. `_segment_by_script` called without the flag must behave exactly as it
   did before this session.
"""

import re

from src.utils.document_generator import DocumentGenerator

HEB_RE = re.compile(r'[֐-׿]')


# ---------------------------------------------------------------------------
# 1. the whole-line verdict
# ---------------------------------------------------------------------------

def test_a_bare_hebrew_line_is_a_hebrew_line():
    assert DocumentGenerator._is_hebrew_only_line("בְּעֶרֶב שַׁבָּת.")


def test_punctuation_and_emphasis_markers_do_not_disqualify_a_line():
    # the piyyut lines arrive fully italicised, and carry their own commas
    assert DocumentGenerator._is_hebrew_only_line("*הָיוּ לְחֶרְפָּה וּלְבִזּוֹת,*")
    # ketiv/qere brackets are the normal shape of a quoted verse line
    assert DocumentGenerator._is_hebrew_only_line("הֵמָּה (יְנוֹעוּן) [יְנִיעוּן] לֶאֱכֹל;")


def test_one_latin_letter_disqualifies_a_line():
    assert not DocumentGenerator._is_hebrew_only_line('Seventeen verses address אֱלֹקִים, "God"')
    assert not DocumentGenerator._is_hebrew_only_line("- B: with rinah (בְּרִנָּה)")


def test_a_line_with_no_hebrew_is_not_a_hebrew_line():
    assert not DocumentGenerator._is_hebrew_only_line("The poet has finished his argument.")
    assert not DocumentGenerator._is_hebrew_only_line("")


# ---------------------------------------------------------------------------
# 2. the segmentation the verdict buys
# ---------------------------------------------------------------------------

def test_quoted_hebrew_line_is_one_rtl_run_including_its_punctuation():
    line = "בְּעֶרֶב שַׁבָּת."
    segs = DocumentGenerator._segment_by_script(line, hebrew_line=True)
    assert segs == [(line, True)], segs


def test_every_terminator_the_guides_actually_use_rides_along():
    for mark in ".,;:!?…—":
        line = "אֲנִי בְצֶדֶק אֶחֱזֶה פָנֶיךָ" + mark
        segs = DocumentGenerator._segment_by_script(line, hebrew_line=True)
        assert segs == [(line, True)], (mark, segs)


def test_surrounding_whitespace_stays_out_of_the_rtl_run():
    segs = DocumentGenerator._segment_by_script("  שָׁלוֹם עֲלֵיכֶם.  ", hebrew_line=True)
    assert segs == [("  ", False), ("שָׁלוֹם עֲלֵיכֶם.", True), ("  ", False)], segs


def test_the_verdict_is_required_the_flag_is_not_a_formality():
    # Same text, no flag: the pre-session behaviour, punctuation in its own LTR run.
    line = "בְּעֶרֶב שַׁבָּת."
    segs = DocumentGenerator._segment_by_script(line)
    assert segs == [("בְּעֶרֶב שַׁבָּת", True), (".", False)], segs


# ---------------------------------------------------------------------------
# 3. mixed prose is untouched
# ---------------------------------------------------------------------------

def test_mixed_prose_keeps_its_punctuation_in_the_english_run():
    text = 'Seventeen verses address אֱלֹקִים, "God" — the generic title'
    segs = DocumentGenerator._segment_by_script(text, hebrew_line=False)
    assert segs == [
        ("Seventeen verses address ", False),
        ("אֱלֹקִים", True),
        (', "God" — the generic title', False),
    ], segs


def test_the_fragment_trap_a_latin_free_fragment_of_an_english_line():
    """` (גּוֹרָל), ` falls out of splitting `the lot *(גּוֹרָל)*, and` on its emphasis
    markers. It is Latin-free, so the verdict MUST come from the caller's line and not
    from the fragment — otherwise the English comma is pulled inside the RTL run and
    renders to the left of the opening parenthesis."""
    fragment = " (גּוֹרָל), "
    assert DocumentGenerator._is_hebrew_only_line(fragment)      # the trap is real
    segs = DocumentGenerator._segment_by_script(fragment)        # ...and the flag is off
    assert (fragment, True) not in segs, segs
    assert any(not is_heb and "," in t for t, is_heb in segs), segs


# ---------------------------------------------------------------------------
# 4. the callers take the verdict at line level, before the markdown split
# ---------------------------------------------------------------------------

def _source_of(func):
    import inspect
    return inspect.getsource(func)


def test_process_markdown_formatting_takes_the_verdict_before_the_emphasis_split():
    src = _source_of(DocumentGenerator._process_markdown_formatting)
    verdict = src.index("_is_hebrew_only_line(modified_text)")
    split = src.index("re.split(")
    assert verdict < split, (
        "_process_markdown_formatting must decide on the whole line BEFORE splitting it "
        "on emphasis markers; deciding per fragment re-opens the ` (גּוֹרָל), ` trap"
    )


def test_soft_break_path_never_asks_the_verdict_of_a_fragment():
    """`_add_paragraph_with_soft_breaks` splits on emphasis BEFORE it splits on newlines,
    so the `line` variable inside its loop is a markdown fragment, not a line. The verdict
    has to be taken once at the top of the method."""
    src = _source_of(DocumentGenerator._add_paragraph_with_soft_breaks)
    assert "_is_hebrew_only_line(modified_text)" in src
    assert "_is_hebrew_only_line(line)" not in src


def test_no_call_site_segments_a_fragment_with_a_hardcoded_true():
    import inspect
    from src.utils import document_generator as module
    src = inspect.getsource(module)
    assert "hebrew_line=True" not in src, (
        "the verdict is always computed from a line, never asserted at a call site"
    )


# ---------------------------------------------------------------------------
# Session 405: the end mark of a Hebrew QUOTATION inside a mixed line
# ---------------------------------------------------------------------------

def _gen_with_paragraph(parts):
    from docx import Document
    gen = DocumentGenerator.__new__(DocumentGenerator)
    gen.document = Document()
    p = gen.document.add_paragraph()
    for text, bold in parts:
        gen._add_inline_runs(p, text, bold=bold)
    gen._join_rtl_runs_across_whitespace()
    gen._attach_hebrew_end_marks()
    return p


def _is_rtl(run):
    from docx.oxml import ns
    rPr = run._element.find(ns.qn('w:rPr'))
    return rPr is not None and rPr.find(ns.qn('w:rtl')) is not None


def test_a_quotations_period_before_its_citation_joins_the_hebrew():
    """Ps 81 arm D, the author: the lines `… תִּשְׁמַע־**לִי**. (v. 9)` "all show periods at the
    RIGHT side of the hebrew". The bold לִי puts the period in its own markdown fragment."""
    p = _gen_with_paragraph([("שְׁמַע עַמִּי וְאָעִידָה בָּךְ, יִשְׂרָאֵל אִם־תִּשְׁמַע־", False),
                             ("לִי", True), (". (v. 9)", False)])
    runs = [r for r in p.runs if r.text]
    i = next(k for k, r in enumerate(runs) if r.text == "לִי")
    assert runs[i + 1].text == "." and _is_rtl(runs[i + 1])
    assert runs[i + 2].text == " (v. 9)" and not _is_rtl(runs[i + 2])
    assert p.text == "שְׁמַע עַמִּי וְאָעִידָה בָּךְ, יִשְׂרָאֵל אִם־תִּשְׁמַע־לִי. (v. 9)"


def test_a_quotations_period_before_its_translation_joins_the_hebrew():
    p = _gen_with_paragraph([('told תִּקְעוּ בַחֹדֶשׁ שׁוֹפָר. "Blow the horn." The', False)])
    runs = [r for r in p.runs if r.text]
    k = next(k for k, r in enumerate(runs) if r.text == ".")
    assert _is_rtl(runs[k]) and runs[k + 1].text.startswith(' "Blow')


def test_a_term_or_a_sentence_going_on_in_english_keeps_its_english_mark():
    for parts in ([('call Him אֱלֹקִים. "God" is', False)],                         # one word
                  [("the line שְׁמַע עַמִּי וְאָעִידָה. Then the poem", False)]):  # English follows
        p = _gen_with_paragraph(parts)
        assert not any(r.text.startswith(".") and _is_rtl(r) for r in p.runs)


def test_a_long_quotations_end_mark_stays_with_its_hebrew_block():
    """Session 405, Ps 81 arm D: a 6+-word quotation is lifted into its own RTL paragraph; its
    period (`X. "Translation."`) opened the NEXT line. It belongs to the Hebrew block."""
    gen = DocumentGenerator.__new__(DocumentGenerator)
    s = ('First comes the prohibition in verse 10: לֹא־יִהְיֶה בְךָ קֵל זָר, וְלֹא תִשְׁתַּחֲוֶה לְקֵל נֵכָר. '
         '"You shall have no foreign god in you." Then')
    before, hebrew, after = gen._split_long_hebrew_block(s)
    assert hebrew.endswith("נֵכָר.") and after.startswith('"You shall')
    comma = s.replace('נֵכָר. "', 'נֵכָר, "')                 # the pre-S405 shape: comma dropped
    before, hebrew, after = gen._split_long_hebrew_block(comma)
    assert hebrew.endswith("נֵכָר") and after.startswith('"You shall')
    bold = ('verse 13: וָאֲשַׁלְּחֵהוּ בִּשְׁרִירוּת לִבָּם, יֵלְכוּ בְּמוֹעֲצוֹתֵי**הֶם**. '
            '"So I let him go." As Meiri')
    before, hebrew, after = gen._split_long_hebrew_block(bold)
    assert hebrew.endswith("בְּמוֹעֲצוֹתֵי**הֶם**.") and after.startswith('"So I let')


def test_a_wrapped_long_quotation_leaves_no_stray_markers_behind():
    """Session 405: `(*X…Y*), celebrating` (the Ps 81 writer's reasoning) printed `*), celebrating`."""
    gen = DocumentGenerator.__new__(DocumentGenerator)
    s = ("verses 4 and 5 of Psalm 81 (*תִּקְעוּ בַחֹדֶשׁ שׁוֹפָר... כִּי חֹק לְיִשְׂרָאֵל הוּא*), "
         "celebrating the biblical mandate")
    before, hebrew, after = gen._split_long_hebrew_block(s)
    assert before == "verses 4 and 5 of Psalm 81"
    assert hebrew.startswith("(תִּקְעוּ") and hebrew.endswith("הוּא)") and "*" not in hebrew
    assert after == "celebrating the biblical mandate"
    open_paren = ('the first of Tishri (אֵיזֶהוּ חַג שֶׁהַחֹדֶשׁ מִתְכַּסֶּה בּוֹ הֱוֵי אוֹמֵר, '
                  '"Which festival is it?"), the poem')
    before, hebrew, after = gen._split_long_hebrew_block(open_paren)
    assert before == "the first of Tishri" and after.startswith('("Which festival is it?"), the poem')
