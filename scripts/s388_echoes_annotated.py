"""
Session 388: the echoes v3.1 dossiers, ANNOTATED (not blind).

For every entry: which model(s) proposed it and why; which judge(s) chose it and at what rank;
both judges' scores (illumination, truth, craft, interest, humour, haunting, memorable,
originality, thought-provoking) and notes; any correction a judge made; whether the text could be
validated (cut from its page at $0); and whether the production dossier -- and guide -- had it.
Then the candidates dropped because another guide already quotes the work, every candidate not
chosen (with both judges' scores), and every production entry with whether v3.1 found it.

    python scripts/s388_echoes_annotated.py 76 77            # reads output/psalm_N/echoes_v3_1/
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from docx.shared import Inches, Pt  # noqa: E402

from s388_echoes_packet import DOCS, Packet, old_entries  # noqa: E402
from src.agents.echoes_v3 import SCORE_KEYS, SCORE_NAMES, same_author, same_work, usable  # noqa: E402
from src.utils.document_generator import export_pdf  # noqa: E402

VERSION_DIR = "echoes_v3_1"
MODEL = {"claude-opus-5-5": "Opus 5.5", "gpt-6-sol": "GPT-6 Sol", "gemini-3.1-pro-preview": "Gemini 3.1 Pro"}
LANE_TITLES = {"literature": "Literary echoes", "beyond": "Resonances beyond literature", "far": "Far associations"}
SHORT = {"ill": "Illum", "tru": "Truth", "cra": "Craft", "int": "Inter", "fun": "Humour", "hau": "Haunt",
         "mem": "Memor", "ori": "Orig", "tho": "Thought"}

TEXT_STATUS = {
    ("literature", "verified"): "Validated: the passage was cut from its source page by the program ($0), so it is "
                                "on that page by construction",
    ("literature", "translation_only"): "Partly validated: the original could not be retrieved; the English "
                                        "translation was cut from its page",
    ("literature", "page_unreadable"): "Not validated: the page was found but refused or could not be read; "
                                       "reference only, no quotation",
    ("literature", "unconfirmed"): "Not validated: the pages found did not yield the passage; reference only",
    ("literature", "not_found"): "Not validated: no page printing the text was found; reference only",
    ("other", "verified"): "Validated: the supporting quotation was found on its source page ($0 check)",
    ("other", "page_unreadable"): "Not validated: gathered from a page that refuses scripts (403), so the program "
                                  "could not confirm it; flagged in the dossier",
    ("other", "unconfirmed"): "Not validated: the gathered quotation was NOT found on its page",
    ("other", "not_found"): "Not validated: no source was found; reference only",
    ("literature", "not_retrieved"): "Not fetched: the dossier did not need this alternate",
    ("other", "not_retrieved"): "Not fetched: the dossier did not need this alternate",
}


def surname(c: dict) -> str:
    words = [w for w in re.findall(r"[^\W\d_]+", c.get("creator") or "") if len(w) >= 4]
    return words[-1] if words else ""


def production_items(psalm: int):
    out = []
    for e in old_entries(psalm):
        m = re.match(r"(.*?),\s*\*(.*?)\*", e["heading"])
        # a heading with no *Work* ("The Baal Cycle (KTU 1.4 V) (c. 1400 BCE)") names a work, not an author
        creator, work = (m.group(1), m.group(2)) if m else ("", re.sub(r"\s*\([^)]*\)", "", e["heading"]))
        out.append({"creator": creator.strip(), "work": work.strip(), "verses": e["verses"]})
    return out


def production_line(c: dict, prod: list, guide: str, lane: str) -> str:
    work = [p for p in prod if same_work(c, p)]
    auth = [p for p in prod if not same_work(c, p) and same_author(c, p)]
    if work:
        s = f"Yes: the production dossier has this work (at {work[0]['verses']})."
    elif auth:
        s = f"Same author, different work: production had *{auth[0]['work']}* (at {auth[0]['verses']})."
    elif lane != "literature":
        s = "No: the production pipeline has no lane for material beyond literature."
    else:
        s = "No: not in the production dossier."
    sn = surname(c)
    if lane == "literature" and sn:
        s += (f" The production guide {'mentions' if re.search(rf'\b{re.escape(sn)}', guide) else 'does not mention'}"
              f" {sn}.")
    return s


def quote_md(e: dict) -> list:
    r = e.get("retrieved") or {}
    if not usable(e):
        return []
    if e["lane"] == "literature":
        src = r.get("original", "") if e.get("status") == "verified" else r.get("translation", "")
        q = [f"> {ln}" if ln.strip() else ">" for ln in src.strip().split("\n")]
        if e.get("status") == "verified" and r.get("check_translation") == "quote found on the page":
            q += [">"] + [f"> {ln}" if ln.strip() else ">" for ln in r["translation"].strip().split("\n")]
        return q + [""]
    return [f"> {ln}" if ln.strip() else ">" for ln in (r.get("quote") or "").strip().split("\n")] + [""]


def total(r: dict) -> int:
    return sum(int(r.get(k, 0) or 0) for k in SCORE_KEYS)


class Annotated(Packet):
    def table(self, rows: list, widths: list, size: float = 8):
        t = self.document.add_table(rows=0, cols=len(widths))
        t.style = "Table Grid"
        t.autofit = False
        for ri, r in enumerate(rows):
            cells = t.add_row().cells
            for ci, w in enumerate(widths):
                cells[ci].width = Inches(w)
                p = cells[ci].paragraphs[0]
                self._process_markdown_formatting(p, str(r[ci]) if ci < len(r) else "", set_font=True)
                for run in p.runs:
                    run.font.size = Pt(size)
                    if ri == 0:
                        run.bold = True
        self.document.add_paragraph()

    def entry(self, e, pool, ratings, prod, guide, lane_counts):
        c = e["candidate"]
        head = f"#### {e['verses']} — {(c.get('creator') + ', ') if c.get('creator') else ''}{c.get('work', '')}"
        head += f" ({c['date']})" if c.get("date") else ""
        if e["lane"] != "literature" and c.get("domain"):
            head += f" — {c['domain']}"
        md = [head.replace("*", ""), ""]
        pos = f"{'Finalist' if e['role'] == 'final' else 'Alternate'} {e['rank']} of {lane_counts[(e['lane'], e['role'])]}"
        if e.get("promoted"):
            pos += ", promoted into the dossier because a finalist's text could not be validated"
        elif e["role"] == "alternate":
            pos += ", kept in reserve"
        by = "; ".join(f"{MODEL.get(m, m)}: {v}" for m, v in e.get("chosen_by", {}).items()) or "neither judge"
        not_by = [MODEL.get(m, m) for m in ratings if m not in e.get("chosen_by", {})]
        md.append(f"**Merged position:** {pos}. **Chosen by:** {by}"
                  f"{'; not chosen by ' + ', '.join(not_by) if not_by else ''}."
                  f"{' Added by the rule that every dossier carries a Jewish or Hebrew poem.' if e.get('jewish_rule') else ''}")
        md.append("")
        md.append(f"**For the writer:** {e.get('move', '')}")
        md.append("")
        if e.get("fix"):
            md.append(f"**Correction from a judge:** {e['fix']}")
            md.append("")
        self.markdown("\n".join(md))
        rows = [["Judge"] + [SHORT[k] for k in SCORE_KEYS] + ["Total /45", "Note"]]
        for m, rs in ratings.items():
            r = rs.get(e["id"])
            if r:
                rows.append([MODEL.get(m, m)] + [r.get(k, "") for k in SCORE_KEYS] + [total(r), r.get("note", "")])
        if len(rows) > 1:
            self.table(rows, [0.85] + [0.42] * 9 + [0.5, 1.9], size=7.5)
        md = []
        same = [p for p in pool if same_work(c, p)] or [c]
        seen = {}
        for p in same:
            seen.setdefault(p["_source"], p)
        md.append("**Proposed by:**")
        for m, p in seen.items():
            flags = [f"confidence {p['confidence']}"] if p.get("confidence") else []
            if p.get("obvious"):
                flags.append("marked 'obvious'")
            loc = f" [{p['locus']}]" if p.get("locus") else ""
            pat = f" *Pattern:* {p['pattern']}" if p.get("pattern") else ""
            md.append(f"- **{MODEL.get(m, m)}**{loc} ({', '.join(flags) or 'no flags'}): {p.get('move', '')}{pat}")
        others = [p for p in pool if not same_work(c, p) and same_author(c, p)]
        if others:
            md.append("- *Same author, other works:* " + "; ".join(
                f"{MODEL.get(p['_source'], p['_source'])}: *{p.get('work', '')}*" for p in others))
        md.append("")
        lk = "literature" if e["lane"] == "literature" else "other"
        r = e.get("retrieved") or {}
        url = r.get("original_url") or r.get("translation_url") or r.get("url") or ""
        md.append(f"**Text:** {TEXT_STATUS.get((lk, e.get('status')), e.get('status', ''))}."
                  f"{' Source: ' + url if url and usable(e) else ''}")
        md.append("")
        md.append(f"**In production?** {production_line(c, prod, guide, e['lane'])}")
        md.append("")
        md += quote_md(e)
        self.markdown("\n".join(md))


def build(psalms):
    out = DOCS / f"Psalms {' and '.join(map(str, psalms))} - Echoes v3.1 annotated (S388).docx"
    g = Annotated(psalms[0], out)
    d = g.document
    d.add_heading("Echoes v3.1, annotated: who proposed what, who chose it, and how it scored", level=1)
    g.note("Session 388, second round. Three models proposed candidates independently (Opus 5.5, GPT-6 Sol, "
           "Gemini 3.1 Pro), in three lanes: literature, resonances beyond literature, and far associations. "
           "Works already quoted in another psalm's guide were removed before judging. Two judges (Opus 5.5 and "
           "Gemini 3.1 Pro) then each scored EVERY candidate 0-5 on nine qualities and chose finalists; their "
           "choices were merged without a further call: picks both made come first, then single picks alternating "
           "by rank. Scores: Illum = illumination, Inter = interest, Memor = memorable, Orig = originality.")
    for psalm in psalms:
        base = ROOT / "output" / f"psalm_{psalm}" / VERSION_DIR
        entries = json.loads((base / "entries.json").read_text(encoding="utf-8"))
        pool = json.loads((base / "pool.json").read_text(encoding="utf-8"))
        rec = json.loads((base / "judge.json").read_text(encoding="utf-8"))
        cost = json.loads((base / "cost_report.json").read_text(encoding="utf-8"))
        prod_cost = json.loads((ROOT / "output" / f"psalm_{psalm}" / "literary_echoes" / "cost_report.json")
                               .read_text(encoding="utf-8"))["total_cost_usd"]
        guide_path = ROOT / "output" / f"psalm_{psalm}" / f"psalm_{psalm:03d}_copy_edited.md"
        guide = guide_path.read_text(encoding="utf-8") if guide_path.exists() else ""
        prod = production_items(psalm)
        ratings = {m: {str(r.get("id")): r for r in j.get("ratings") or []} for m, j in rec["judges"].items()}
        lane_counts = {}
        for e in entries:
            lane_counts[(e["lane"], e["role"])] = lane_counts.get((e["lane"], e["role"]), 0) + 1

        d.add_page_break()
        d.add_heading(f"Psalm {psalm}", level=2)
        per_model = ", ".join(f"{MODEL.get(m, m)} {n}" for m, n in cost["pool"].items())
        stage_rows = [["Stage", "Model", "Cost"]]
        for k, v in cost["stages"].items():
            stage_rows.append([k.split(" [")[0], MODEL.get(k.split("[")[-1].rstrip("]"), k.split("[")[-1].rstrip("]")),
                               f"${v['usd']:.2f}"])
        g.note(f"Proposals: {per_model}. {len(rec.get('dropped_as_used') or [])} dropped as already quoted in "
               f"another guide; {len(pool)} judged. In the dossier: "
               + ", ".join(f"{cost['used'].get(l, 0)} {LANE_TITLES[l].lower()}" for l in LANE_TITLES)
               + f". Cost ${cost['total_usd']:.2f} (production pipeline: ${prod_cost:.2f}).")
        g.table(stage_rows, [2.5, 1.6, 0.8])
        for m, j in rec["judges"].items():
            if j.get("notes"):
                g.note(f"{MODEL.get(m, m)}'s note on the pool: {j['notes']}")

        for lane, title in LANE_TITLES.items():
            d.add_heading(f"{title} in the dossier", level=3)
            used = sorted([e for e in entries if e["lane"] == lane and e.get("use")],
                          key=lambda e: int(m.group(1)) if (m := re.search(r":(\d+)", e["verses"])) else 999)
            for e in used:
                g.entry(e, pool, ratings, prod, guide, lane_counts)

        spare = [e for e in entries if e["role"] == "alternate" and not e.get("use")]
        if spare:
            d.add_heading("Alternates the dossier did not need", level=3)
            for e in spare:
                g.entry(e, pool, ratings, prod, guide, lane_counts)

        dropped = rec.get("dropped_as_used") or []
        if dropped:
            d.add_heading(f"Dropped before judging: already quoted in another guide ({len(dropped)})", level=3)
            rows = [["Proposed", "By", "Already quoted in"]]
            for p in dropped:
                rows.append([f"{(p.get('creator') + ', ') if p.get('creator') else ''}*{p.get('work', '')}*",
                             MODEL.get(p["_source"], p["_source"]), f"Ps {p.get('used_in')}"])
            g.table(rows, [4.2, 1.2, 1.4])

        chosen = [e["candidate"] for e in entries]
        rest = [p for p in pool if not any(same_work(p, c) for c in chosen)]
        d.add_heading(f"Everything else proposed ({len(rest)} candidates neither judge kept)", level=3)
        g.note("Duplicates of a kept work are folded into its entry above. Scores are Illumination / Truth / "
               "Total of nine (max 45).")
        names = list(ratings)
        rows = [["Lane", "Verses", "Proposed", "By", "Proposer's reason"] +
                [f"{MODEL.get(m, m)}: score, note" for m in names]]
        lane_order = {l: i for i, l in enumerate(LANE_TITLES)}
        for p in sorted(rest, key=lambda p: (lane_order.get(p.get("lane"), 9),
                                             int(m.group(1)) if (m := re.search(r":(\d+)", p.get("verses", ""))) else 999)):
            name = f"{(p.get('creator') + ', ') if p.get('creator') else ''}*{p.get('work', '')}*"
            cells = []
            for m in names:
                r = ratings[m].get(p["id"])
                cells.append(f"{r.get('ill', '')}/{r.get('tru', '')}/{total(r)}: {r.get('note', '')}"
                             + (f" FIX: {r['fix']}" if r.get("fix") else "") if r else "—")
            rows.append([p.get("lane", "")[:4], p.get("verses", ""), name, MODEL.get(p["_source"], p["_source"]),
                         p.get("move", "")] + cells)
        g.table(rows, [0.4, 0.45, 1.15, 0.55, 1.75] + [1.25] * len(names), size=7)

        d.add_heading("The production dossier: did v3.1 find the same things?", level=3)
        rows = [["Verses", "Production entry", "In the v3.1 pool?", "In the dossier?"]]
        dropped_all = pool + dropped
        for q in prod:
            hits = [p for p in dropped_all if same_work(p, q)]
            ahits = [p for p in dropped_all if not same_work(p, q) and same_author(p, q)]
            kept = [e for e in entries if e.get("use") and same_work(e["candidate"], q)]
            was_dropped = any(same_work(p, q) for p in dropped)
            pooled = (", ".join(sorted({MODEL.get(p["_source"], p["_source"]) for p in hits}))
                      + (" (dropped: used elsewhere)" if was_dropped else "")) if hits else \
                (f"same author, other work ({', '.join(sorted({MODEL.get(p['_source'], p['_source']) for p in ahits}))})"
                 if ahits else "no")
            rows.append([q["verses"], f"{q['creator'] + ', ' if q['creator'] else ''}*{q['work']}*", pooled,
                         "yes" if kept else "no"])
        g.table(rows, [0.6, 3.2, 2.1, 0.9])

    g.finish()
    export_pdf(out)
    return out


if __name__ == "__main__":
    build([int(a) for a in sys.argv[1:]] or [76, 77])
