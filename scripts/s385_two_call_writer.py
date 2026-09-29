"""
Session 385 — the two-call "forest" writer on Psalm 76 (experimental; archive after the session).

  Call 1  the INTRODUCTION ESSAY, under the S384 P1 ("forest") instructions, byte-identical
          to the trials.
  Call 2  the rest of the guide (liturgical section, verse-by-verse commentary, reader
          questions), as a second turn of the same conversation: the essay is the assistant
          turn, the new instructions (VERSE_INSTRUCTIONS below) are the next user turn.

Every call sends the same first user turn, [B's inputs block][P1 instructions], with a
5-minute cache_control on its last block, so call 2 reads what call 1 wrote. A cache entry
lives 5 minutes from the START of the request that wrote or read it; call 1 took 84-114 s in
the S384 trials. If call 1 is still generating at KEEPALIVE_AFTER_S, one max_tokens=0 request
refreshes the entry (a cache read, ~$0.04). A miss is not a failure: call 2 then pays a fresh
write (~$0.9).

Two variants of call 2, run in parallel once call 1 is done:
  new  on the essay call 1 just wrote. Its thinking blocks are replayed unchanged, so
       Opus 5.5 sees its own essay reasoning (the connections it listed and set aside).
  F    on essay F from the S384 trials, the author's pick. Text only: F's thinking was not
       kept in replayable form.

Inputs (all in the repo):
  archive/psalm_76_S384_essay_trials/prompts/inputs_block.txt    B's writer inputs, byte for byte
  archive/psalm_76_S384_essay_trials/prompts/p1_instructions.txt  the P1 text the trials sent
  archive/psalm_76_S384_essay_trials/essays/P1_opus_high_r1/raw_response.md   essay F

Writes ONLY under archive/psalm_76_S385_two_call/. Resumable: a finished call is never re-run.
Hard spend cap: no new call starts once the ledger passes BUDGET_USD.

    python scripts/s385_two_call_writer.py --dry-run        # build + check the prompts; $0, no key
    python scripts/s385_two_call_writer.py                  # run (needs ANTHROPIC_API_KEY)
    python scripts/s385_two_call_writer.py --docx           # rebuild the reading DOCXs from saved output; $0
    python scripts/s385_two_call_writer.py --selftest-docx  # render the S383 all-5.5 guide through the
                                                            # same DOCX path, to test it; $0
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

PSALM = 76
MODEL = "claude-opus-5-5"
EFFORT = "high"            # both calls: effort is part of the cached prefix, so it must match
THINKING = {"type": "adaptive", "display": "summarized"}
BUDGET_USD = 8.0
KEEPALIVE_AFTER_S = 240

TRIALS = ROOT / "archive" / "psalm_76_S384_essay_trials"
INPUTS_FILE = TRIALS / "prompts" / "inputs_block.txt"
P1_FILE = TRIALS / "prompts" / "p1_instructions.txt"
ESSAY_F_FILE = TRIALS / "essays" / "P1_opus_high_r1" / "raw_response.md"
B_DIR = ROOT / "archive" / "psalm_76_S383_opus55_ab" / "B_opus55"
STATS_FILE = B_DIR / "psalm_076_pipeline_stats.json"      # same dossier; used for the DOCX methods page
OUT = ROOT / "archive" / "psalm_76_S385_two_call"
LEDGER_PATH = OUT / "ledger.json"

LIT_MARKER = "---LITURGICAL-SECTION-START---"


# ---------------------------------------------------------------------------
# The verse-commentary instructions (call 2's user turn)
# ---------------------------------------------------------------------------

VERSE_INSTRUCTIONS = """## ═══════════════════════════════════════════════════════════════════════════
## NEXT: THE REST OF THE GUIDE — LITURGY, VERSE COMMENTARY, QUESTIONS
## ═══════════════════════════════════════════════════════════════════════════

Your introduction essay is finished. It will be printed first, exactly as you wrote it, and the reader will have just read it. Now write the rest of the study guide: a short section on the psalm in Jewish liturgy, the verse-by-verse commentary, and a few questions for the reader.

Everything in the instructions you were given for the essay still holds: who you are writing as, READ LIKE A POET, BRING THE WHOLE LIBRARY and its two tests, COMMENTATORS, LITERARY ECHOES, and WRITING (including "you are the author" — never mention the research or anything behind it). What follows is what changes when you move from the essay to the verses.

## WHAT THE VERSE COMMENTARY IS FOR

The essay made one argument about the whole poem. The commentary walks the reader through the poem slowly, a verse at a time, and stops wherever there is something worth seeing or hearing that the essay did not show. It is where the reader, text in hand, finds that each line holds more than they thought: the word that is stranger than its translation admits, the sound that carries a meaning, the line a later reader heard differently, the older text this one answers, the rabbinic argument built on a single letter.

What the essay said is spent. Do not re-argue it, re-quote its evidence, or reach its conclusions by a second route. Where a verse carried the essay's argument, point back to it in one sentence at most ("the introduction took up this line's missing 'us'") and spend the note on what the essay could not use. The best material for the notes is what you gathered for the essay and left out: the connections you listed and set aside, the commentator who saw something off the essay's line, the variant, the liturgical use, the pattern in the concordance.

## HOW TO WORK (your reasoning phase)

1. **Take stock of the essay.** Note briefly what it established and which evidence it used. That material is spent.
2. **Walk the poem.** For each verse, before you consult the research: what would a careful listener stumble on here? What does this line do that prose would not?
3. **Search the research, verse by verse.** Read the commentators on the verse, its verse notes and lexicon entries, the concordance and figurative-language material, the Greek, the liturgical uses, and any cross-verse observations or shared-vocabulary parallels that touch it. Put the same question to your own knowledge: what does this line open onto?
4. **Choose, and choose unevenly.** One or two things per verse: the ones that change how the line reads. A verse holding a real discovery gets a long note; a verse of routine construction may need three sentences. If your notes come out about the same length, they were filled, not written.

## WHAT A NOTE CAN HOLD (a menu, not a checklist)

- **What the line does**: a sound, a grammatical swerve, a proportion, a word held back or placed last, and what it does to a listener.
- **A strange or rare word**: what it means, where else it occurs (quote the other passage, Hebrew and English), and what the comparison shows. "Occurs only here" or "only twice" is a claim about the whole Bible; make it only when the concordance material in front of you supports it.
- **A crux**: when the verse admits two readings, name them, then either choose one and say why, or show why the ambiguity works.
- **A commentator who sees something**: a reading that changes the line, a disagreement, a risk. Work from the commentator's own text in front of you, not from memory, and keep the shape of his argument: which is his reading, which is his proof, and which is his alternative. Show the step that takes him from the words to the reading.
- **The rabbinic afterlife**: where the Talmud or midrash (often through the Torah Temimah) mined the verse, and the move that got them there.
- **The Greek**: when the Septuagint read the Hebrew differently and the difference teaches something about the text or its first readers. Ration it: at most two verses in five.
- **The psalm in use**: a prayer that quotes the line (quote the prayer in Hebrew and English, name the service and the rite, and say whether it follows the plain sense or puts the words to new use), or a documented moment when someone used these words.
- **An echo**: see below.

## LITERARY ECHOES ACROSS THE WHOLE GUIDE

Aim for about three echoes for every four verses across the whole guide, counting those already in your essay: for this psalm's {n_verses} verses, about {echo_target} in all. That is a level to reach with good echoes, not a quota. Rank the research's echoes together with any you know that fit better, and spend from the top. An echo earns its place by showing something about THIS line: the same event told by another voice, the same image turned to the opposite use, the same problem solved another way. A poem that shares only a mood is decoration. Do not repeat an echo the essay has used. Quote fully, and set poems lineated, as the essay instructions describe.

## ACCURACY

You are drawing on your own knowledge as well as the research, so hold every claim to one test: state only what you know to be accurate, name sources precisely, and quote from memory only what you are sure of. When you are unsure of a detail, say less. When your memory and a text in front of you disagree about what that text says, the text wins.

## THE FORMAT OF EACH VERSE

Cover every verse, in order, including the superscription (verse 1):

1. A header on its own line: `**Verse N**`. You may group two to four closely linked verses under one header (`**Verses 5–6**`), but each verse still gets its own Hebrew line and its own translation, and all of them come before the shared note.
2. The Hebrew of the verse, punctuated to show its poetic structure: semicolons or commas between the cola, a full stop at the end.
3. Your own English translation of the whole verse, as a one-line block quote beginning `> `. Every word; no ellipsis, brackets, alternatives, Hebrew or comment. Where a word is genuinely undecidable, translate the reading you argue for and let the note do the arguing.
4. The note, in paragraphs. Prose, not lists.

The translation line has rendered every word, so nothing in the note is owed to completeness. A phrase you pass over is not a gap.

Use the supplied phonetic transcriptions only when a sound carries the point: in backticks, with the capitalized stress left as given; when you claim that two words share a sound, bold the letters that carry it.

## THE LITURGICAL SECTION (200–500 words)

Open your response with the exact marker `---LITURGICAL-SECTION-START---` on its own line. Then use `####` subheadings: `#### Full psalm` (where and when the whole psalm is recited) and `#### Key verses` (verses or phrases quoted in prayers). Quote the Hebrew of both the psalm and the prayer, with English. Say what the placement reveals, and whether the liturgy follows the plain sense or puts the words to new use. If the research contains Shimush Tehillim material, add `#### Practical Kabbalah`: state the prescribed use, and suggest briefly what in the psalm's language makes the association intelligible; omit the subsection if there is none. Every specific liturgical use in the research should appear somewhere in the guide, here or in a verse note. In this section, refer to verses in running text ("v. 3"); never begin a line with **Verse.

The rites, which are easy to confuse: **Nusach Ashkenaz** is the rite of non-Hasidic Ashkenazi Jews. **Nusach Sefard** is the HASIDIC rite, used by Ashkenazi Hasidim; never call it "Sephardic." **Edot HaMizrach** is the rite of the Sephardic and Middle Eastern communities. When the research says "Sefard," it means the Hasidic rite.

## QUESTIONS FOR THE READER

End with 4–6 questions to be printed before the commentary: specific, answerable from the guide, each pointing toward something the essay or the notes will show.

## OUTPUT

Return exactly this shape, with nothing before the marker:

---LITURGICAL-SECTION-START---

#### Full psalm
...

#### Key verses
...

### VERSE COMMENTARY

**Verse 1**

[Hebrew]

> [translation]

[note]

**Verse 2**
...

### REFINED READER QUESTIONS

1. ...
2. ...
"""


# ---------------------------------------------------------------------------
# Prompt assembly
# ---------------------------------------------------------------------------

def psalm_verses(inputs: str) -> dict:
    """{verse: {'hebrew', 'english'}} from the inputs block's PSALM TEXT section."""
    a = inputs.index("### PSALM TEXT")
    b = inputs.index("### STRUCTURAL OVERVIEW")
    out = {}
    for m in re.finditer(r"### Verse (\d+)\n\*\*Hebrew:\*\* (.*?)\n\*\*English:\*\* (.*?)\n\*\*Phonetic:\*\*",
                         inputs[a:b], re.S):
        out[int(m.group(1))] = {"hebrew": m.group(2).strip(), "english": " ".join(m.group(3).split())}
    return out


def build() -> dict:
    inputs = INPUTS_FILE.read_text(encoding="utf-8")
    p1 = P1_FILE.read_text(encoding="utf-8")
    essay_f = ESSAY_F_FILE.read_text(encoding="utf-8")
    verses = psalm_verses(inputs)
    n = len(verses)
    verse_instr = (VERSE_INSTRUCTIONS.replace("{n_verses}", str(n))
                   .replace("{echo_target}", str(round(n * 0.75))))
    assert "{n_verses}" not in verse_instr and "{echo_target}" not in verse_instr
    first_turn = [
        {"type": "text", "text": inputs},
        {"type": "text", "text": p1, "cache_control": {"type": "ephemeral"}},  # 5-minute TTL
    ]
    return {"inputs": inputs, "p1": p1, "essay_f": essay_f, "verses": verses, "n_verses": n,
            "verse_instr": verse_instr, "first_turn": first_turn}


def first_turn_messages(P) -> list:
    return [{"role": "user", "content": P["first_turn"]}]


def verse_call_messages(P, essay_content: list) -> list:
    return first_turn_messages(P) + [
        {"role": "assistant", "content": essay_content},
        {"role": "user", "content": [{"type": "text", "text": P["verse_instr"]}]},
    ]


# ---------------------------------------------------------------------------
# Ledger and cost
# ---------------------------------------------------------------------------

_ledger_lock = threading.Lock()


def _load_ledger() -> dict:
    if LEDGER_PATH.exists():
        return json.loads(LEDGER_PATH.read_text(encoding="utf-8"))
    return {"calls": []}


def spent() -> float:
    return round(sum(c["cost_usd"] for c in _load_ledger()["calls"]), 4)


def _record(entry: dict) -> None:
    with _ledger_lock:
        led = _load_ledger()
        led["calls"].append(entry)
        OUT.mkdir(parents=True, exist_ok=True)
        LEDGER_PATH.write_text(json.dumps(led, ensure_ascii=False, indent=2), encoding="utf-8")


def _usage(us) -> dict:
    cc = getattr(us, "cache_creation", None)
    w5 = (getattr(cc, "ephemeral_5m_input_tokens", 0) or 0) if cc else (getattr(us, "cache_creation_input_tokens", 0) or 0)
    w1h = (getattr(cc, "ephemeral_1h_input_tokens", 0) or 0) if cc else 0
    return {"input": us.input_tokens or 0, "cache_write": w5, "cache_write_1h": w1h,
            "cache_read": getattr(us, "cache_read_input_tokens", 0) or 0,
            "output": us.output_tokens or 0}


def _cost(u: dict) -> float:
    from src.utils.cost_tracker import resolve_pricing
    p = resolve_pricing(MODEL)
    return (u["input"] * p["input"] + u["cache_write"] * p["cache_write"]
            + u["cache_write_1h"] * p["cache_write_1h"] + u["cache_read"] * p["cache_read"]
            + u["output"] * p["output"]) / 1_000_000


class OverBudget(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# API calls
# ---------------------------------------------------------------------------

_client = None


def client():
    """PSALMS_ANTHROPIC_API_KEY first: in a Claude Code cloud environment a plain
    ANTHROPIC_API_KEY would also be visible to Claude Code itself, so the key is stored
    under a name only this project reads. Falls back to the SDK default (ANTHROPIC_API_KEY,
    e.g. from the author's local .env)."""
    global _client
    if _client is None:
        import os

        import anthropic
        try:  # local runs: the project's .env supplies ANTHROPIC_API_KEY
            from dotenv import load_dotenv
            load_dotenv(ROOT / ".env")
        except ImportError:
            pass
        key = os.environ.get("PSALMS_ANTHROPIC_API_KEY")
        kw = {"api_key": key} if key else {}
        _client = anthropic.Anthropic(max_retries=4, timeout=3600.0, **kw)
    return _client


def keepalive(P, label: str) -> None:
    """Refresh the 5-minute entry: same prefix, thinking and effort; max_tokens=0, no stream."""
    try:
        t0 = time.time()
        msg = client().messages.create(model=MODEL, max_tokens=0, thinking=THINKING,
                                       output_config={"effort": EFFORT},
                                       messages=first_turn_messages(P))
        u = _usage(msg.usage)
        _record({"label": label, "model": MODEL, "effort": EFFORT, "usage": u,
                 "cost_usd": round(_cost(u), 4), "seconds": round(time.time() - t0),
                 "stop_reason": msg.stop_reason})
        print(f"  [{label}] cache refreshed: read {u['cache_read']:,}, wrote {u['cache_write']:,}", flush=True)
    except Exception as e:  # a failed keep-alive costs at most a cache miss later
        print(f"  [{label}] keep-alive failed ({str(e)[:160]}); call 2 may pay a fresh cache write", flush=True)


def stream_call(label: str, messages: list, P=None, keepalive_after: float | None = None) -> dict:
    if spent() >= BUDGET_USD:
        raise OverBudget(f"budget ${BUDGET_USD} reached before '{label}' (spent ${spent()})")
    last_err = None
    for attempt in range(3):
        timer = None
        if keepalive_after and P is not None:
            timer = threading.Timer(keepalive_after, keepalive, args=(P, f"{label}_keepalive"))
            timer.daemon = True
            timer.start()
        try:
            t0 = time.time()
            with client().messages.stream(model=MODEL, max_tokens=128000, thinking=THINKING,
                                          output_config={"effort": EFFORT},
                                          messages=messages) as stream:
                msg = stream.get_final_message()
            u = _usage(msg.usage)
            cost = _cost(u)
            _record({"label": label, "model": MODEL, "effort": EFFORT, "usage": u,
                     "cost_usd": round(cost, 4), "seconds": round(time.time() - t0),
                     "stop_reason": msg.stop_reason})
            text = "".join(b.text for b in msg.content if b.type == "text")
            thinking = "\n\n".join(b.thinking for b in msg.content if b.type == "thinking" and b.thinking)
            if msg.stop_reason == "refusal":
                raise RuntimeError(f"refusal: {getattr(msg, 'stop_details', None)}")
            if msg.stop_reason != "end_turn" or not text.strip():
                raise RuntimeError(f"stop_reason={msg.stop_reason}, text chars={len(text)}")
            print(f"  [{label}] {round(time.time() - t0)}s, ${cost:.4f}; in {u['input']:,} "
                  f"+ cache read {u['cache_read']:,} + cache write {u['cache_write']:,}; out {u['output']:,}",
                  flush=True)
            return {"text": text, "thinking": thinking, "usage": u, "cost": cost,
                    "content": [b.to_dict() for b in msg.content]}
        except OverBudget:
            raise
        except Exception as e:
            last_err = e
            print(f"  [{label}] attempt {attempt + 1} failed: {str(e)[:200]}", flush=True)
            time.sleep(15 * (attempt + 1))
        finally:
            if timer is not None:
                timer.cancel()
    raise RuntimeError(f"{label}: all attempts failed: {last_err}")


def replayable(content: list) -> list:
    """Assistant content to send back unchanged: thinking blocks byte-for-byte with their
    signatures (Opus 5.5 reads its own earlier reasoning), text blocks as text."""
    out = []
    for b in content:
        t = b.get("type")
        if t == "thinking":
            out.append({"type": "thinking", "thinking": b["thinking"], "signature": b["signature"]})
        elif t == "redacted_thinking":
            out.append({"type": "redacted_thinking", "data": b["data"]})
        elif t == "text":
            out.append({"type": "text", "text": b["text"]})
    return out


# ---------------------------------------------------------------------------
# Output checks and assembly
# ---------------------------------------------------------------------------

VERSE_HEAD = re.compile(r"(?m)^\*\*Verses?\s+(\d+)(?:\s*[–-]\s*(\d+))?\*\*\s*$")


def check_structure(rest: str, n_verses: int) -> list:
    """Problems with call 2's output that would break the guide downstream."""
    problems = []
    if not rest.lstrip().startswith(LIT_MARKER):
        problems.append("response does not open with the liturgical marker")
    if not re.search(r"(?m)^#{1,4}\s*VERSE COMMENTARY\s*$", rest):
        problems.append("no VERSE COMMENTARY header")
        return problems
    lit = rest[:re.search(r"(?m)^#{1,4}\s*VERSE COMMENTARY\s*$", rest).start()]
    if re.search(r"(?m)^\*\*Verses?\s+\d", lit):
        problems.append("a line in the liturgical section begins with **Verse (breaks the copy editor's split)")
    heads = list(VERSE_HEAD.finditer(rest))
    covered = []
    for i, h in enumerate(heads):
        lo, hi = int(h.group(1)), int(h.group(2) or h.group(1))
        covered += list(range(lo, hi + 1))
        body = rest[h.end(): heads[i + 1].start() if i + 1 < len(heads) else len(rest)]
        n_trans = len(re.findall(r"(?m)^> \S", body))
        if n_trans < hi - lo + 1:
            problems.append(f"verse(s) {lo}-{hi}: fewer '> ' lines than verses")
    if sorted(covered) != list(range(1, n_verses + 1)):
        problems.append(f"verses covered {sorted(covered)} != 1..{n_verses}")
    m = re.search(r"(?ms)^#{1,4}\s*REFINED READER QUESTIONS\s*$(.*)", rest)
    if not m:
        problems.append("no REFINED READER QUESTIONS")
    elif len(re.findall(r"(?m)^\s*\d+\.\s+\S", m.group(1))) < 4:
        problems.append("fewer than 4 reader questions")
    return problems


def parse_guide(full: str) -> dict:
    """The production parser's regexes (MasterEditorV2._parse_writer_response)."""
    res = {"introduction": "", "verse_commentary": "", "reader_questions": ""}
    vm = re.search(r"#{1,4}\s*VERSE COMMENTARY\s*\n(.*?)(?=#{1,4}\s*REFINED READER QUESTIONS|$)", full, re.S | re.I)
    if vm:
        res["verse_commentary"] = vm.group(1).strip()
    im = re.search(r"#{1,4}\s*INTRODUCTION ESSAY\s*\n(.*?)(?=#{1,4}\s*VERSE COMMENTARY|$)", full, re.S | re.I)
    if im:
        res["introduction"] = im.group(1).strip()
    qm = re.search(r"#{1,4}\s*REFINED READER QUESTIONS\s*\n(.*?)$", full, re.S | re.I)
    if qm:
        res["reader_questions"] = qm.group(1).strip()
    return res


def write_guide_files(vdir: Path, essay_text: str, rest: str) -> dict:
    full = essay_text.strip() + "\n\n" + rest.strip() + "\n"
    (vdir / "full_writer_response.md").write_text(full, encoding="utf-8")
    g = parse_guide(full)
    (vdir / f"psalm_{PSALM:03d}_edited_intro.md").write_text(g["introduction"], encoding="utf-8")
    (vdir / f"psalm_{PSALM:03d}_edited_verses.md").write_text(g["verse_commentary"], encoding="utf-8")
    qs = [m.group(1).strip() for m in re.finditer(r"(?m)^\s*\d+\.\s+(.+)$", g["reader_questions"])]
    (vdir / f"psalm_{PSALM:03d}_reader_questions.json").write_text(
        json.dumps({"psalm_number": PSALM, "curated_questions": qs, "source": "s385_two_call"},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    return g


# ---------------------------------------------------------------------------
# Metrics (descriptive, for comparison with the S383 all-5.5 guide K)
# ---------------------------------------------------------------------------

COMMENTATORS = ["Rashi", "Ibn Ezra", "Radak", "Meiri", "Alshich", "Romemot El", "Minchat Shai", "Norzi",
                "Metzudat", "Chida", "Chomat Anakh", "Malbim", "Torah Temimah", "Talmud", "Midrash"]


def verse_metrics(verse_md: str) -> dict:
    heads = list(VERSE_HEAD.finditer(verse_md))
    per = []
    for i, h in enumerate(heads):
        body = verse_md[h.end(): heads[i + 1].start() if i + 1 < len(heads) else len(verse_md)]
        prose = "\n".join(l for l in body.splitlines() if not l.startswith(">"))
        label = "v" + h.group(1) + (f"-{h.group(2)}" if h.group(2) else "")
        per.append({"verses": label, "words": len(prose.split()),
                    "greek": bool(re.search(r"Septuagint|LXX|Greek", body)),
                    "block_quote_lines": sum(1 for l in body.splitlines() if l.startswith(">"))})
    return {
        "words_total": len(verse_md.split()),
        "per_verse": per,
        "commentator_mentions": {c: len(re.findall(re.escape(c), verse_md)) for c in COMMENTATORS
                                 if re.search(re.escape(c), verse_md)},
        "verses_with_greek": f"{sum(p['greek'] for p in per)} of {len(per)}",
    }


def shared_ngrams(a: str, b: str, n: int = 6) -> int:
    def grams(t):
        w = re.findall(r"[\w֐-׿']+", t.lower())
        return {tuple(w[i:i + n]) for i in range(len(w) - n + 1)}
    return len(grams(a) & grams(b))


def report() -> None:
    rows = {}
    k_full = (B_DIR / "master_writer_v4_response_psalm_76.txt").read_text(encoding="utf-8")
    kg = parse_guide(k_full)
    k_essay = kg["introduction"].split(LIT_MARKER)[0]
    rows["K (S383 one call)"] = (k_essay, kg)
    for v in ("new", "F"):
        f = OUT / v / "full_writer_response.md"
        if f.exists():
            g = parse_guide(f.read_text(encoding="utf-8"))
            rows[f"two-call {v}"] = (g["introduction"].split(LIT_MARKER)[0], g)
    out = {}
    for name, (essay, g) in rows.items():
        m = verse_metrics(g["verse_commentary"])
        m["essay_words"] = len(essay.split())
        m["liturgy_words"] = len(g["introduction"].split(LIT_MARKER)[1].split()) if LIT_MARKER in g["introduction"] else 0
        m["essay_verse_shared_6grams"] = shared_ngrams(essay, g["verse_commentary"])
        out[name] = m
        print(f"\n== {name}: essay {m['essay_words']:,} words, liturgy {m['liturgy_words']:,}, "
              f"verses {m['words_total']:,}; Greek in {m['verses_with_greek']} notes; "
              f"6-grams shared essay/verses {m['essay_verse_shared_6grams']}")
        print("   per note:", ", ".join(f"{p['verses']}={p['words']}" for p in m["per_verse"]))
        print("   commentators:", m["commentator_mentions"])
    (OUT / "metrics.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# Reading DOCX: production DocumentGenerator + S383's compact layout + thinking appendix
# ---------------------------------------------------------------------------

def build_docx(vdir: Path, title_note: str, thinking_parts: list, out_path: Path) -> Path:
    """Render vdir's intro/verses/questions through the production DocumentGenerator, in the
    S383 compact print layout, with the writer's summarized thinking appended. The psalm
    text comes from the inputs block, so tanakh.db (not in the repo) is not needed."""
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    import src.utils.document_generator as dg

    verses = psalm_verses(INPUTS_FILE.read_text(encoding="utf-8"))

    class _V:
        def __init__(self, n, d):
            self.verse, self.hebrew, self.english = n, d["hebrew"], d["english"]

    class _StubDB:  # stands in for TanakhDatabase().get_psalm()
        def get_psalm(self, n):
            return type("P", (), {"verses": [_V(k, verses[k]) for k in sorted(verses)]})()

    body_pt, appx_pt = 10.5, 9
    scale = body_pt / 12

    def cs_size(style, pt):
        rpr = style.element.get_or_add_rPr()
        el = rpr.find(qn("w:szCs"))
        if el is None:
            el = OxmlElement("w:szCs")
            rpr.append(el)
        el.set(qn("w:val"), str(int(round(pt * 2))))

    def hairline(style, color="A6A6A6"):
        ppr = style.element.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", color)):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.append(bdr)

    class Compact(dg.DocumentGenerator):
        def _set_default_styles(self):
            super()._set_default_styles()
            st = self.document.styles
            st["Normal"].font.size = Pt(body_pt)
            cs_size(st["Normal"], body_pt + 1)
            st["Normal"].paragraph_format.space_after = Pt(4)
            st["BodySans"].font.size = Pt(body_pt)
            cs_size(st["BodySans"], body_pt + 1)
            st["SummaryText"].font.size = Pt(8.5)
            cs_size(st["SummaryText"], 9.5)
            for name, size, before, after in (("Heading 1", 16, 0, 6), ("Heading 2", 12.5, 10, 4),
                                              ("Heading 3", 11, 7, 1), ("Heading 4", body_pt, 6, 1)):
                h = st[name]
                h.font.size = Pt(size)
                h.paragraph_format.space_before = Pt(before)
                h.paragraph_format.space_after = Pt(after)
                h.paragraph_format.keep_with_next = True
            hairline(st["Heading 2"])
            appx = st.add_style("AppendixText", 1)
            appx.base_style = st["BodySans"]
            appx.font.size = Pt(appx_pt)
            cs_size(appx, appx_pt + 1)
            appx.paragraph_format.space_after = Pt(3)
            for s in self.document.sections:
                s.left_margin = s.right_margin = Inches(0.75)
                s.top_margin = Inches(0.7)
                s.bottom_margin = Inches(0.65)
                s.footer_distance = Inches(0.35)

        # Last hook before the complex-script fix-ups and save(), as in S383.
        def _join_rtl_runs_across_whitespace(self):
            self._compact_body()
            self._append_thinking()
            super()._join_rtl_runs_across_whitespace()

        def _compact_body(self):
            body = self.document.element.body
            for p in list(body.iter(qn("w:p"))):
                brs = [br for br in p.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]
                if not brs:
                    continue
                if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip():
                    for br in brs:
                        br.getparent().remove(br)
                    continue
                prev = p.getprevious()
                if prev is not None and prev.tag == qn("w:p") and not "".join(
                        t.text or "" for t in prev.iter(qn("w:t"))).strip():
                    body.remove(prev)
                body.remove(p)
            for tag in ("w:sz", "w:szCs"):
                for el in body.iter(qn(tag)):
                    el.set(qn("w:val"), str(max(15, int(round(int(el.get(qn("w:val"))) * scale)))))
            for sp in body.iter(qn("w:spacing")):
                for k in ("w:before", "w:after"):
                    if sp.get(qn(k)) is not None:
                        sp.set(qn(k), str(int(sp.get(qn(k))) // 2))
            for ind in body.iter(qn("w:ind")):
                for k in ("w:left", "w:start", "w:right", "w:end"):
                    if ind.get(qn(k)) == "720":
                        ind.set(qn(k), "504")
            for tbl in body.iter(qn("w:tbl")):
                for gc in tbl.iter(qn("w:gridCol")):
                    gc.set(qn("w:w"), "5040")
                for tcw in tbl.iter(qn("w:tcW")):
                    tcw.set(qn("w:type"), "dxa")
                    tcw.set(qn("w:w"), "5040")

        def _append_thinking(self):
            doc = self.document
            doc.add_heading("Appendix: The Writer's Reasoning", level=2)
            note = doc.add_paragraph(style="AppendixText")
            r = note.add_run(title_note)
            r.italic = True
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            two = doc.add_section(WD_SECTION.CONTINUOUS)
            cols = two._sectPr.find(qn("w:cols"))
            if cols is None:
                cols = OxmlElement("w:cols")
                two._sectPr.append(cols)
            cols.set(qn("w:num"), "2")
            cols.set(qn("w:space"), "360")
            start = len(doc.element.body)
            for heading, text in thinking_parts:
                if not text.strip():
                    continue
                self._add_paragraph_with_markdown(f"**{heading}**", style="AppendixText")
                for para in [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]:
                    self._add_paragraph_with_markdown(" ".join(para.split("\n")), style="AppendixText")
            for el in list(doc.element.body)[start:]:
                for sz in el.iter(qn("w:sz")):
                    sz.set(qn("w:val"), str(appx_pt * 2))
                for sz in el.iter(qn("w:szCs")):
                    sz.set(qn("w:val"), str((appx_pt + 1) * 2))
            end = doc.add_section(WD_SECTION.CONTINUOUS)
            end_cols = end._sectPr.find(qn("w:cols"))
            if end_cols is not None:
                end_cols.set(qn("w:num"), "1")

    saved_db = dg.TanakhDatabase
    dg.TanakhDatabase = _StubDB
    try:
        qfile = vdir / f"psalm_{PSALM:03d}_reader_questions.json"
        Compact(PSALM, vdir / f"psalm_{PSALM:03d}_edited_intro.md", vdir / f"psalm_{PSALM:03d}_edited_verses.md",
                STATS_FILE, out_path, qfile if qfile.exists() else None).generate()
    finally:
        dg.TanakhDatabase = saved_db
    print(f"  DOCX: {out_path.relative_to(ROOT)}")
    return out_path


THINKING_NOTE = ("The summarized reasoning Claude Opus 5.5 returned while writing this guide (the API's "
                 "thinking display, not a verbatim trace), in two parts: the essay call and the call that "
                 "wrote the liturgy, verse commentary and questions. This guide has NOT been copy edited.")


def docx_for_variant(v: str) -> Path | None:
    vdir = OUT / v
    if not (vdir / "full_writer_response.md").exists():
        return None
    parts = []
    ess_think = (OUT / "essay_new" / "thinking.txt") if v == "new" else None
    if ess_think and ess_think.exists():
        parts.append(("Essay call", ess_think.read_text(encoding="utf-8")))
    elif v == "F":
        f_think = TRIALS / "essays" / "P1_opus_high_r1" / "thinking_summary.txt"
        if f_think.exists():
            parts.append(("Essay call (essay F, Session 384 trials)", f_think.read_text(encoding="utf-8")))
    vt = vdir / "thinking.txt"
    if vt.exists():
        parts.append(("Verse-commentary call", vt.read_text(encoding="utf-8")))
    label = "new essay" if v == "new" else "essay F"
    return build_docx(vdir, THINKING_NOTE, parts, vdir / f"Psalm 76 - two-call forest ({label}).docx")


def selftest_docx() -> None:
    """Render the S383 all-5.5 guide (K) through build_docx: exercises parse -> files -> DOCX at $0."""
    vdir = OUT / "_selftest_K"
    vdir.mkdir(parents=True, exist_ok=True)
    full = (B_DIR / "master_writer_v4_response_psalm_76.txt").read_text(encoding="utf-8")
    essay, rest = full.split(LIT_MARKER, 1)
    g = write_guide_files(vdir, essay, LIT_MARKER + rest)
    print(f"  parsed: intro {len(g['introduction']):,} chars, verses {len(g['verse_commentary']):,}, "
          f"questions {len(g['reader_questions']):,}")
    print("  structure problems (K, written under the one-call prompt):",
          check_structure(LIT_MARKER + rest, 13) or "none")
    think = (B_DIR / "psalm_076_master_writer_v4_thinking.txt").read_text(encoding="utf-8")
    build_docx(vdir, "Self-test render of the S383 guide.", [("One-call writer", think)],
               vdir / "selftest_K.docx")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def run(P) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "prompts").mkdir(exist_ok=True)
    (OUT / "prompts" / "verse_instructions.txt").write_text(P["verse_instr"], encoding="utf-8")

    # Call 1: the essay (writes the 5-minute cache entry on the first turn).
    e_dir = OUT / "essay_new"
    if not (e_dir / "content.json").exists():
        print("[call 1] essay ...", flush=True)
        r = stream_call("essay_new", first_turn_messages(P), P=P, keepalive_after=KEEPALIVE_AFTER_S)
        e_dir.mkdir(parents=True, exist_ok=True)
        (e_dir / "content.json").write_text(json.dumps(r["content"], ensure_ascii=False, indent=1), encoding="utf-8")
        (e_dir / "essay.md").write_text(r["text"], encoding="utf-8")
        (e_dir / "thinking.txt").write_text(r["thinking"], encoding="utf-8")
    else:
        print("[call 1] essay already written; call 2 may miss the cache (fresh write)", flush=True)
    essay_content = replayable(json.loads((e_dir / "content.json").read_text(encoding="utf-8")))
    essay_new_text = (e_dir / "essay.md").read_text(encoding="utf-8")

    # Call 2, both variants in parallel (each reads the cached first turn).
    variants = {
        "new": (essay_new_text, essay_content),
        "F": (P["essay_f"], [{"type": "text", "text": P["essay_f"]}]),
    }

    def one(v):
        vdir = OUT / v
        if (vdir / "call2_response.md").exists():
            return v, "already done"
        try:  # one variant failing must not lose the other
            r = stream_call(f"verses_{v}", verse_call_messages(P, variants[v][1]))
        except Exception as e:
            return v, f"FAILED: {str(e)[:200]}"
        vdir.mkdir(parents=True, exist_ok=True)
        (vdir / "call2_response.md").write_text(r["text"], encoding="utf-8")
        (vdir / "thinking.txt").write_text(r["thinking"], encoding="utf-8")
        return v, "done"

    print("[call 2] verse commentary on the new essay and on essay F ...", flush=True)
    with ThreadPoolExecutor(2) as ex:
        for v, status in ex.map(one, list(variants)):
            print(f"  [{v}] {status}", flush=True)

    for v, (essay_text, _) in variants.items():
        vdir = OUT / v
        if not (vdir / "call2_response.md").exists():
            continue
        rest = (vdir / "call2_response.md").read_text(encoding="utf-8")
        problems = check_structure(rest, P["n_verses"])
        (vdir / "structure_check.json").write_text(json.dumps(problems, indent=2), encoding="utf-8")
        print(f"  [{v}] structure: {problems or 'OK'}", flush=True)
        write_guide_files(vdir, essay_text, rest)
        try:
            docx_for_variant(v)
        except Exception as e:
            print(f"  [{v}] DOCX failed: {e}", flush=True)
    report()
    print(f"\nTotal spent: ${spent():.4f} (ledger: {LEDGER_PATH.relative_to(ROOT)})")


def dry_run(P) -> None:
    print(f"inputs block: {len(P['inputs']):,} chars; P1: {len(P['p1']):,}; "
          f"verse instructions: {len(P['verse_instr']):,}; essay F: {len(P['essay_f']):,}")
    print(f"verses parsed from inputs: {P['n_verses']} (echo target {round(P['n_verses'] * 0.75)})")
    assert P["n_verses"] == 13 and P["essay_f"].lstrip().startswith("### INTRODUCTION ESSAY")
    m1 = first_turn_messages(P)
    m2 = verse_call_messages(P, [{"type": "text", "text": P["essay_f"]}])
    assert m2[0] == m1[0], "call 2 must open with call 1's exact first turn (cache prefix)"
    print("cache prefix: call 2's first turn is identical to call 1's (OK)")
    # Token estimate from the S384 trials: the same first turn measured 185,599 Opus tokens.
    est_in = 181_208 + 4_391
    from src.utils.cost_tracker import resolve_pricing
    p = resolve_pricing(MODEL)
    c1 = (est_in * p["cache_write"] + 10_000 * p["output"]) / 1e6
    c2 = (est_in * p["cache_read"] + 7_000 * p["input"] + 45_000 * p["output"]) / 1e6
    print(f"estimate: call 1 ≈ ${c1:.2f}; each call 2 ≈ ${c2:.2f}; total ≈ ${c1 + 2 * c2:.2f}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--docx", action="store_true")
    ap.add_argument("--selftest-docx", action="store_true")
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args()
    P = build()
    if a.dry_run:
        return dry_run(P)
    if a.selftest_docx:
        return selftest_docx()
    if a.docx:
        for v in ("new", "F"):
            docx_for_variant(v)
        return
    if a.report:
        return report()
    run(P)


if __name__ == "__main__":
    main()
