"""Session 383: rebuild the Ps 76 arm-B DOCX in a compact print layout, with the
Master Writer's summarized thinking appended as a two-column appendix.

Goes through the production DocumentGenerator (same inputs the pipeline's STEP 6
used), so Hebrew/RTL handling, divine-name modification and the quote-block
rendering are the real ones. Only the page design is changed, via a subclass:

  * Letter, margins 0.75" sides / 0.7" top / 0.65" bottom
  * body 10.5pt (Hebrew 11.5pt, keeping the house +1pt complex-script ratio)
  * paragraph spacing halved; no forced page breaks
  * section heads with a hairline rule beneath; verse heads tight to their text
  * appendix: 9pt in two balanced columns

Nothing in src/ is modified.
"""
import os
import re
import sys
from pathlib import Path

REPO = Path(r"C:\dev\personal\psalms")
os.chdir(REPO)
sys.path.insert(0, str(REPO))
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

from docx.enum.section import WD_SECTION  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

from src.utils.document_generator import DocumentGenerator  # noqa: E402

B_DIR = REPO / "output" / "psalm_76" / "_opus55_B"
THINKING = B_DIR / "psalm_076_master_writer_v4_thinking.txt"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "Documents" / "Psalm study guide" / "Psalm 76 (B - Opus 5.5).docx"

BODY_PT = 10.5
APPX_PT = 9
SCALE = BODY_PT / 12  # run-level sizes were authored against a 12pt body


def _set_cs_size(style, pt):
    rpr = style.element.get_or_add_rPr()
    el = rpr.find(qn("w:szCs"))
    if el is None:
        el = OxmlElement("w:szCs")
        rpr.append(el)
    el.set(qn("w:val"), str(int(round(pt * 2))))


def _hairline_below(style, color="A6A6A6"):
    ppr = style.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


class CompactGenerator(DocumentGenerator):
    def _set_default_styles(self):
        super()._set_default_styles()
        st = self.document.styles

        st["Normal"].font.size = Pt(BODY_PT)
        _set_cs_size(st["Normal"], BODY_PT + 1)
        st["Normal"].paragraph_format.space_after = Pt(4)
        st["BodySans"].font.size = Pt(BODY_PT)
        _set_cs_size(st["BodySans"], BODY_PT + 1)
        st["SummaryText"].font.size = Pt(8.5)
        _set_cs_size(st["SummaryText"], 9.5)

        for name, size, before, after in (
            ("Heading 1", 16, 0, 6),
            ("Heading 2", 12.5, 10, 4),
            ("Heading 3", 11, 7, 1),
            ("Heading 4", BODY_PT, 6, 1),
        ):
            h = st[name]
            h.font.size = Pt(size)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        _hairline_below(st["Heading 2"])

        appx = st.add_style("AppendixText", 1)
        appx.base_style = st["BodySans"]
        appx.font.size = Pt(APPX_PT)
        _set_cs_size(appx, APPX_PT + 1)
        appx.paragraph_format.space_after = Pt(3)

        for s in self.document.sections:
            s.left_margin = s.right_margin = Inches(0.75)
            s.top_margin = Inches(0.7)
            s.bottom_margin = Inches(0.65)
            s.footer_distance = Inches(0.35)

    # Last hook before the complex-script fix-ups and save(): do the layout
    # post-processing and append the appendix here, so the fix-ups cover it too.
    def _join_rtl_runs_across_whitespace(self):
        self._compact_body()
        self._append_thinking()
        super()._join_rtl_runs_across_whitespace()

    def _compact_body(self):
        body = self.document.element.body

        # 1. Drop forced page breaks (and the empty carrier paragraph before each).
        for p in list(body.iter(qn("w:p"))):
            brs = [br for br in p.iter(qn("w:br")) if br.get(qn("w:type")) == "page"]
            if not brs:
                continue
            text = "".join(t.text or "" for t in p.iter(qn("w:t")))
            if text.strip():
                for br in brs:
                    br.getparent().remove(br)
                continue
            prev = p.getprevious()
            if prev is not None and prev.tag == qn("w:p") and not "".join(
                t.text or "" for t in prev.iter(qn("w:t"))
            ).strip():
                body.remove(prev)
            body.remove(p)

        # 2. Scale explicit run sizes authored for a 12pt body.
        for tag in ("w:sz", "w:szCs"):
            for el in body.iter(qn(tag)):
                v = int(el.get(qn("w:val")))
                el.set(qn("w:val"), str(max(15, int(round(v * SCALE)))))

        # 3. Halve explicit paragraph spacing; pull quote indents in.
        for sp in body.iter(qn("w:spacing")):
            for k in ("w:before", "w:after"):
                if sp.get(qn(k)) is not None:
                    sp.set(qn(k), str(int(sp.get(qn(k))) // 2))
        for ind in body.iter(qn("w:ind")):
            for k in ("w:left", "w:start", "w:right", "w:end"):
                if ind.get(qn(k)) == "720":  # the 0.5" quote indent
                    ind.set(qn(k), "504")    # 0.35"

        # 4. Psalm table: span the wider text block.
        for tbl in body.iter(qn("w:tbl")):
            for gc in tbl.iter(qn("w:gridCol")):
                gc.set(qn("w:w"), "5040")
            for tcw in tbl.iter(qn("w:tcW")):
                tcw.set(qn("w:type"), "dxa")
                tcw.set(qn("w:w"), "5040")

    def _append_thinking(self):
        doc = self.document
        doc.add_heading("Appendix: The Master Writer's Reasoning", level=2)
        note = doc.add_paragraph(style="AppendixText")
        r = note.add_run(
            "The summarized reasoning Claude Opus 5.5 returned while drafting this guide "
            "(the API's thinking display, not a verbatim trace). It was written before the "
            "copy edit, so a detail it weighs may differ from the final text."
        )
        r.italic = True
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

        two_col = doc.add_section(WD_SECTION.CONTINUOUS)
        cols = two_col._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            two_col._sectPr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")  # 0.25" gutter

        start = len(doc.element.body)
        paras = [p.strip() for p in re.split(r"\n\s*\n", THINKING.read_text(encoding="utf-8")) if p.strip()]
        for para in paras:
            self._add_paragraph_with_markdown(" ".join(para.split("\n")), style="AppendixText")

        # Hebrew blocks inside the helper hard-code 13pt; hold the appendix to its own size.
        for el in list(doc.element.body)[start:]:
            for sz in el.iter(qn("w:sz")):
                sz.set(qn("w:val"), str(APPX_PT * 2))
            for sz in el.iter(qn("w:szCs")):
                sz.set(qn("w:val"), str((APPX_PT + 1) * 2))

        # A closing continuous break makes Word balance the two columns on the last page.
        end = doc.add_section(WD_SECTION.CONTINUOUS)
        end_cols = end._sectPr.find(qn("w:cols"))
        if end_cols is not None:
            end_cols.set(qn("w:num"), "1")
        print(f"Appendix: {len(paras)} paragraphs from {THINKING.name}")


gen = CompactGenerator(
    76,
    B_DIR / "psalm_076_edited_intro.md",
    B_DIR / "psalm_076_edited_verses.md",
    B_DIR / "psalm_076_pipeline_stats.json",
    OUT,
    None,
)
gen.generate()
