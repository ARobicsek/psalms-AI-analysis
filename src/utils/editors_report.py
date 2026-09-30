"""
The editors' report (Session 387): what the fact checker, the citation verifier and the
copy editor did to one psalm's guide, and what every stage of the run cost.

Built from files the pipeline already writes, so it costs $0 and can be rebuilt at any time:

    psalm_NNN_fact_check.json            every claim, its verdict, evidence, stage, checker
    psalm_NNN_fact_check_telemetry.json  every lookup and web action, the gathered web
                                         passages and their page checks, the evidence
                                         handed over up front, per-call usage
    psalm_NNN_copy_edit_changes.md       the copy editor's numbered change log
    psalm_NNN_copy_edited.md             the text after the copy edit
    psalm_NNN_citation_verification.md   the $0 SQL citation check, before the copy edit
    psalm_NNN_post_copy_edit_citations.md  citation mismatches the copy edit introduced
    psalm_NNN_cost.json                  per-model and per-stage cost, non-token charges
    psalm_NNN_writer_calls.json          the forest writer's calls

It is descriptive: it reports what happened and leaves the judging to the author.

    python -m src.utils.editors_report 77        # writes output/psalm_77/psalm_077_editors_report.{md,docx}
"""

from __future__ import annotations

import difflib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# ---------------------------------------------------------------------------
# Parsing (pure)
# ---------------------------------------------------------------------------

_ENTRY = re.compile(r"(?ms)^(\d+)\.\s+(.*?)(?=^\d+\.\s|^###\s|\Z)")
_CATS = re.compile(r"^\[([^\]]+)\]\s*")

# Copy-editor categories (COPY_EDITOR_SYSTEM_PROMPT). 7 and 9(g) are the factual ones.
CATEGORY_NAMES = {
    "1": "structural claims", "2": "internal inconsistency", "3": "form/content confusion",
    "4": "negative citation", "5": "Hebrew script", "6": "weak cross-cultural parallel",
    "7": "factual and textual accuracy", "8": "Hebrew grammar bloat", "9": "strained argument",
    "10": "unexplained technical term", "11": "banned phrase",
}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFC", s or "")
    s = re.sub(r"[*_`]", "", s)
    s = s.translate(str.maketrans({"“": '"', "”": '"', "‘": "'", "’": "'"}))
    return re.sub(r"\s+", " ", s).strip()


def parse_change_log(text: str) -> Tuple[List[Dict], List[str]]:
    """(entries, unverified notes) from a copy_edit_changes.md."""
    body = text
    unverified: List[str] = []
    m = re.search(r"(?m)^#{2,4}\s*UNVERIFIED\s*$", text)
    if m:
        body, tail = text[:m.start()], text[m.end():]
        for line in tail.splitlines():
            line = line.strip()
            if line.startswith("#"):
                break
            if line and line not in ("-", "*"):
                unverified.append(re.sub(r"^[-*\d.)\s]+", "", line).strip())
        unverified = [u for u in unverified if u]
    entries = []
    for em in _ENTRY.finditer(body):
        raw = em.group(2).strip()
        # Leading tags in any order: named ones ([FACT-CHECK], [CITATION FIX]) are flags,
        # numeric ones ([7], [9b, 10]) are categories.
        rest, cats, flags = raw, [], []
        while (cm := _CATS.match(rest)):
            parts = [c.strip() for c in cm.group(1).split(",") if c.strip()]
            if parts and all(re.match(r"^\d", c) for c in parts):
                cats += parts
            else:
                flags.append(cm.group(1).strip().upper())
            rest = rest[cm.end():]
        tagged = "FACT-CHECK" in flags
        lm = re.match(r"\*\*(.+?)\*\*:?\s*", rest)
        location = lm.group(1).strip() if lm else ""
        if lm:
            rest = rest[lm.end():]
        what, _, why = rest.partition("\nWhy:")
        if not why:
            what, _, why = rest.partition("Why:")
        if not why and "\n" in rest.strip():   # rationale on its own indented line, no "Why:"
            what, _, why = rest.strip().partition("\n")
        entries.append({"n": int(em.group(1)), "categories": cats, "location": location,
                        "what": " ".join(what.split()), "why": " ".join(why.split()),
                        "fact_check_tagged": tagged, "citation_fix": "CITATION FIX" in flags})
    return entries, unverified


def is_factual(entry: Dict) -> bool:
    return any(c == "7" or c.replace(" ", "").lower() in ("9g", "9(g)") for c in entry["categories"])


def _sentences(text: str) -> List[str]:
    out = []
    for para in re.split(r"\n\s*\n", text):
        para = " ".join(para.split())
        out += [s for s in re.split(r"(?<=[.!?])\s+(?=[\"“(A-Z֐-׿])", para) if len(s) > 20]
    return out


def outcome_for(record: Dict, final_text: str, entries: List[Dict]) -> Dict:
    """Did the copy editor act on a contradicted claim? Exact whole-sentence test, then the
    closest sentence in the final text, then the [FACT-CHECK] log entry that best matches."""
    sent = norm(record.get("sentence", ""))
    final_n = norm(final_text)
    unchanged = bool(sent) and sent in final_n
    now = ""
    if not unchanged and sent:
        best, score = "", 0.0
        for s in _sentences(final_text):
            r = difflib.SequenceMatcher(None, sent, norm(s)).ratio()
            if r > score:
                best, score = s, r
        if score >= 0.6:  # below this the closest sentence is usually unrelated
            now = best
    words = set(re.findall(r"\w{4,}", (record.get("claim", "") + " " + record.get("suggested_fix", "") or "").lower()))
    match, mscore = None, 0
    loc = (record.get("location") or "").lower()
    for e in entries:
        text = (e["what"] + " " + e["why"]).lower()
        sc = len(words & set(re.findall(r"\w{4,}", text)))
        if e["location"] and loc and (e["location"].lower() in loc or loc in e["location"].lower()):
            sc += 3
        if e["fact_check_tagged"]:
            sc += 1
        if sc > mscore:
            match, mscore = e, sc
    if mscore < 5:
        match = None
    if unchanged:
        status = "NOT changed: the sentence is still in the final text word for word"
    elif match and match["fact_check_tagged"]:
        status = f"changed; logged as [FACT-CHECK] change #{match['n']}"
    elif match:
        status = f"changed; the closest log entry (#{match['n']}) is NOT tagged [FACT-CHECK]"
    else:
        status = "changed, with no matching change-log entry"
    return {"unchanged": unchanged, "now_reads": now, "log_entry": match, "status": status}


def parse_citation_report(text: str) -> List[Dict]:
    out = []
    for m in re.finditer(r"(?ms)^### Issue \d+: (.+?)\n(.*?)(?=^### |\Z)", text or ""):
        body = m.group(2)

        def field(name):
            fm = re.search(rf"- \*\*{name}\*\*: (.+)", body)
            return fm.group(1).strip() if fm else ""
        out.append({"ref": m.group(1).strip(), "location": field("Location"),
                    "type": field("Type").strip("`"), "quoted": field("Quoted"),
                    "issue": field("Likely issue")})
    return out


# ---------------------------------------------------------------------------
# Formatting helpers (pure)
# ---------------------------------------------------------------------------

# Plain names for the fact checker's lookup tools (underscores would render as italics).
TOOL_NAMES = {"get_verse": "verse", "get_commentary": "commentator", "search_tanakh": "Bible word search",
              "get_text": "Sefaria text", "search_research": "research search"}


def _usd(x: float) -> str:
    return f"${x:,.2f}" if x >= 0.995 else f"${x:,.3f}"


def _k(n: int) -> str:
    return f"{n:,}" if n else "–"


def _cell(s) -> str:
    return str(s if s is not None else "").replace("|", "/").replace("\n", " ")


def _table(header: List[str], rows: List[List]) -> List[str]:
    out = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    out += ["| " + " | ".join(_cell(c) for c in r) + " |" for r in rows]
    return out + [""]


def _clip(s: str, n: int) -> str:
    s = " ".join((s or "").split())
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


# ---------------------------------------------------------------------------
# The report
# ---------------------------------------------------------------------------

def _load(out: Path, psalm: int) -> Dict:
    stem = f"psalm_{psalm:03d}"

    def rd(name):
        f = out / f"{stem}_{name}"
        return f.read_text(encoding="utf-8") if f.exists() else ""

    def js(name):
        t = rd(name)
        return json.loads(t) if t.strip() else {}

    # Session 388: a fact check OLDER than this guide's print-ready text belongs to an earlier
    # version of the guide (a later run without --fact-check leaves the old file in place), and
    # reporting it would pair another guide's 316 claims with this guide's copy edit. The
    # print-ready file is written in STEP 5, before any fact check (STEP 5a3/4) of the same run.
    fc_file, anchor = out / f"{stem}_fact_check.json", out / f"{stem}_print_ready.md"
    stale = fc_file.exists() and anchor.exists() and fc_file.stat().st_mtime < anchor.stat().st_mtime
    if stale:
        return {"fc": {}, "tele": {}, "stale_fact_check": True,
                "changes": rd("copy_edit_changes.md"), "final": rd("copy_edited.md"),
                "cit_pre": rd("citation_verification.md"), "cit_post": rd("post_copy_edit_citations.md"),
                "cost": js("cost.json"), "writer": js("writer_calls.json"), "stats": js("pipeline_stats.json"),
                "print_ready": rd("print_ready.md")}
    return {"fc": js("fact_check.json"), "tele": js("fact_check_telemetry.json"),
            "changes": rd("copy_edit_changes.md"), "final": rd("copy_edited.md"),
            "cit_pre": rd("citation_verification.md"), "cit_post": rd("post_copy_edit_citations.md"),
            "cost": js("cost.json"), "writer": js("writer_calls.json"), "stats": js("pipeline_stats.json"),
            "print_ready": rd("print_ready.md")}


def build_markdown(psalm: int, out_dir: Path) -> str:
    D = _load(out_dir, psalm)
    fc, tele, cost = D["fc"], D["tele"], D["cost"]
    records = fc.get("claims", [])
    # Reports written before Session 387 can carry six-word stubs on contradicted claims.
    from src.agents.fact_checker import expand_sentence
    for r in records:
        if r.get("verdict") != "supported" and D["print_ready"]:
            full = expand_sentence(r.get("sentence", ""), D["print_ready"])
            if full != r.get("sentence"):
                r["sentence"], r["sentence_expanded"] = full, True
    meta = fc.get("meta", {})
    entries, unverified = parse_change_log(D["changes"])
    by_v = Counter(r["verdict"] for r in records)
    by_stage = Counter(r.get("stage", "local") for r in records)
    contra = [r for r in records if r["verdict"] == "contradicted"]
    outcomes = {r["id"]: outcome_for(r, D["final"], entries) for r in contra}
    n_unchanged = sum(o["unchanged"] for o in outcomes.values())
    n_tagged_match = sum(bool(o["log_entry"] and o["log_entry"]["fact_check_tagged"]) and not o["unchanged"]
                         for o in outcomes.values())
    tagged = [e for e in entries if e["fact_check_tagged"]]
    untagged_fact = [e for e in entries if is_factual(e) and not e["fact_check_tagged"]]
    cit_pre = parse_citation_report(D["cit_pre"])
    cit_post = parse_citation_report(D["cit_post"])
    trace = tele.get("trace", [])
    lookups = [t for t in trace if t.get("kind") == "lookup"]
    web_actions = [t for t in trace if t.get("kind") == "web"]
    gathered = tele.get("gathered", [])
    n_sources = sum(len(g.get("sources", [])) for g in gathered)
    n_verified = sum(1 for g in gathered for s in g.get("sources", []) if s.get("verified"))
    total_cost = cost.get("total_cost", 0.0)
    stages = cost.get("stages", [])
    # Session 390: a resumed run CONTINUES the cost file (stages carry `attempt`), so the file's
    # total is every run of this psalm ($13.77 on Ps 77, whose guide run cost $4.11). Report this
    # run (the latest attempt) and, when there were earlier ones, all runs beside it. The fact
    # check / copy editor rows take the LATEST such stage, not the first.
    last_attempt = max([s.get("attempt", 1) for s in stages] or [1])
    run_cost = sum(s["cost_usd"] for s in stages if s.get("attempt", 1) == last_attempt)
    has_fc = bool(fc)   # False on a run without --fact-check, or with only a stale one (see _load)
    fc_stage = next((s for s in reversed(stages) if s["stage"] == "fact check"), None) if has_fc else None
    ce_stage = next((s for s in reversed(stages) if s["stage"] == "copy editor"), None)
    ce_model = (D["stats"].get("model_usage") or {}).get("copy_editor") or "gpt-5.4"

    L: List[str] = [f"# Psalm {psalm}: What the Editors Did", ""]
    if has_fc:
        L += [f"This report follows the guide to Psalm {psalm} from the moment the writer finished it. Three "
              "things checked or changed it: a **citation verifier** that compares every quoted Bible verse with "
              "the Masoretic text; a **fact checker** that lists every checkable claim in the guide and rules on "
              "each with evidence; and a **copy editor** that received both reports and edited the text. The "
              "report says what each one did. It does not grade them; that is the reader's job. The cost of the "
              "run comes after the summary.", ""]
    else:
        L += [f"This report follows the guide to Psalm {psalm} from the moment the writer finished it. Two things "
              "checked or changed it: a **citation verifier** that compares every quoted Bible verse with the "
              "Masoretic text, and a **copy editor** that received its report and edited the text. **No fact "
              "check ran on this version of the guide** (the pipeline's fact check was off)"
              + (", and the fact check saved beside it belongs to an earlier version, so it is not reported"
                 if D.get("stale_fact_check") else "")
              + ". The copy editor therefore had no evidence to correct facts from. The report says what each "
              "one did. It does not grade them; that is the reader's job. The cost of the run comes after the "
              "summary.", ""]

    # --- at a glance -----------------------------------------------------------
    L += ["## At a glance", ""]
    fc_rows = [
        ["Claims the fact checker listed and ruled on", f"{len(records)}"],
        ["… supported / contradicted / unverifiable",
         f"{by_v.get('supported', 0)} / {by_v.get('contradicted', 0)} / {by_v.get('unverifiable', 0)}"],
        ["… settled from local evidence (no web) / sent to the web",
         f"{by_stage.get('local', 0) + by_stage.get('review', 0)} / {by_stage.get('web', 0)}"],
        ["Contradicted claims the copy editor changed / left word for word",
         f"{len(contra) - n_unchanged} / {n_unchanged}"],
        ["… of the changed ones, matched to a [FACT-CHECK] log entry", f"{n_tagged_match}"],
    ]
    ce_rows = [["Copy-editor changes in all", f"{len(entries)}"]]
    if has_fc:
        ce_rows += [["… tagged [FACT-CHECK]", f"{len(tagged)}"],
                    ["… factual (category 7 or 9g) but NOT tagged [FACT-CHECK]", f"{len(untagged_fact)}"]]
    else:
        ce_rows += [["… factual (category 7 or 9g), made without a fact-check report", f"{len(untagged_fact)}"]]
    ce_rows += [
        ["UNVERIFIED notes the copy editor left for the author", f"{len(unverified)}"],
        ["Quoted-verse mismatches found before the copy edit (after the false-positive filter)", f"{len(cit_pre)}"],
    ]
    if has_fc:   # the post-copy-edit re-check runs only with the fact check (STEP 5b½)
        ce_rows += [["Quoted-verse mismatches the copy edit introduced", f"{len(cit_post)}"]]
    web_rows = [
        ["Free lookups the fact checker made (verse, commentator, concordance, Sefaria, research)",
         f"{len(lookups)}"],
        ["Web searches", f"{meta.get('web_searches', 0)}"],
        ["Web passages gathered / found on their live page", f"{n_sources} / {n_verified}"],
    ]
    if last_attempt > 1:
        cost_rows = [[f"Cost of this run (run {last_attempt} in this psalm's cost file)", _usd(run_cost)],
                     [f"… all {last_attempt} recorded runs of this psalm together", _usd(total_cost)]]
    else:
        cost_rows = [["Cost of this run", _usd(total_cost)]]
    ce_usd = _usd(ce_stage["cost_usd"]) if ce_stage else "–"
    if has_fc:
        cost_rows += [["… fact check (incl. web-search fees) / copy editor",
                       f"{_usd(fc_stage['cost_usd']) if fc_stage else '–'} / {ce_usd}"]]
    else:
        cost_rows += [["… copy editor", ce_usd]]
    rows = (fc_rows if has_fc else []) + ce_rows + (web_rows if has_fc else []) + cost_rows
    L += _table(["", "Count"], rows)

    # --- who is who -------------------------------------------------------------
    models = tele.get("models", {})
    lm, wm, jm = models.get("local", ["?", "?"]), models.get("web_gather", ["?", "?"]), models.get("web_judge", ["?", "?"])
    L += ["## The editors, and what each could see", ""]
    L += ["- **Citation verifier** ($0 SQL check, then a model filter): every Hebrew quotation the guide "
          "attributes to a Bible verse is compared with that verse in the Masoretic text; mismatches go to a "
          "model that throws out false alarms (ellipses, divine-name spellings). Survivors go to the copy editor."]
    if not has_fc:
        L += [f"- **Copy editor** ({ce_model}): gets the guide and the citation report. With no fact-check "
              "report it has no licence to correct facts; doubts about them go in an UNVERIFIED list for the "
              "author, not into the text.", ""]
    else:
        L += [
          f"- **Fact checker, stage 1** ({lm[0]}, reasoning effort {lm[1]}, no web): reads the guide in "
          f"{len(meta.get('chunks', []))} chunks. Each chunk comes with the commentators' entries on its verses "
          "and the text of every verse it cites, handed over free. It can look things up at no cost: a verse, a "
          "commentator's entry, a Hebrew word across the Bible, any Sefaria text (Talmud, midrash, siddur), and "
          "the writer's research. It rules supported / contradicted / unverifiable, or hands a claim on as "
          "needing the web.",
          f"- **Fact checker, stage 2** (web): {wm[0]} (effort {wm[1]}) searches the web and gathers passages "
          "for each handed-on claim, without ruling. Every passage is then fetched from its page by code, free, "
          f"to confirm it is really there. {jm[0]} (effort {jm[1]}) rules from the passages; only a passage "
          "found on its page can contradict the guide.",
          f"- **Copy editor** ({ce_model}): gets the guide, the citation report and the fact-check report. It "
          "may correct a fact only where the report says contradicted, tagging the change [FACT-CHECK]; doubts "
          "about anything else go in an UNVERIFIED list for the author, not into the text.", ""]

    # --- cost ---------------------------------------------------------------------
    L += _cost_section(D, fc_stage)

    if has_fc:   # Session 390: fact-check sections only when a fact check ran
        # --- contradicted -----------------------------------------------------------------
        L += ["## Claims the fact checker contradicted", ""]
        if not contra:
            L += ["None.", ""]
        for r in contra:
            o = outcomes[r["id"]]
            L += [f"### {r['id']} · {r.get('location', '')} · {r.get('claim_type', '')}", ""]
            L += [f"**The guide said:** {r.get('sentence', '')}", "",
                  f"**The claim checked:** {r.get('claim', '')}", "",
                  f"**Finding:** {r.get('explanation', '')}", ""]
            for e in r.get("evidence", []):
                src = e.get("source", "") + (f" ({e['url']})" if e.get("url") else "")
                L += [f"- *Evidence, {src}:* {e.get('quote', '')}"]
            L += [""]
            if r.get("suggested_fix"):
                L += [f"**Suggested fix:** {r['suggested_fix']}", ""]
            L += [f"**Ruled by:** {r.get('checked_by', '')}" + (" (stage 1 first said contradicted)"
                                                                 if r.get("first_verdict") else ""), "",
                  f"**Copy editor:** {o['status']}.", ""]
            if o["now_reads"]:
                L += [f"**The final text now reads:** {o['now_reads']}", ""]
            if o["log_entry"]:
                e = o["log_entry"]
                L += [f"**Its log entry #{e['n']}:** {e['what']}" + (f" *Why:* {e['why']}" if e["why"] else ""), ""]

        # --- web -------------------------------------------------------------------------
        # Keyed by claim: the gathered item keeps stage 1's sentence, which may be a six-word stub
        # while the record's sentence has been expanded.
        web_recs = {norm(r.get("claim", "")): r for r in records if r.get("stage") == "web"}
        L += ["## Claims sent to the web", ""]
        L += [f"Stage 1 handed {len(gathered)} claim(s) on as needing the web. For each: what the gatherer found, "
              "whether code found the passage on its live page, and the judge's verdict.", ""]
        for g in gathered:
            rec = web_recs.get(norm(g.get("claim", "")), {})
            L += [f"### {rec.get('id', '?')} · {g.get('location', '')} · verdict: {rec.get('verdict', '?')}", "",
                  f"**Claim:** {g.get('claim', '')}", ""]
            if not g.get("sources"):
                L += ["- The gatherer returned no passage." + (f" Its note: {g['researcher_note']}"
                                                              if g.get("researcher_note") else "")]
            for s in g.get("sources", []):
                mark = "on its page ✓" if s.get("verified") else f"NOT confirmed ({s.get('check', '')})"
                L += [f"- *{s.get('title') or s.get('source', '')}* ({s.get('url', '')}), {mark}: "
                      f"{_clip(s.get('quote', ''), 400)}"]
            if g.get("sources") and g.get("researcher_note"):
                L += [f"- Gatherer's note: {g['researcher_note']}"]
            if rec.get("explanation"):
                L += ["", f"**Judge:** {rec['explanation']}"]
            L += [""]

        # --- unverifiable -------------------------------------------------------------------
        unv = [r for r in records if r["verdict"] == "unverifiable"]
        L += ["## Claims no evidence could settle (unverifiable)", ""]
        L += _table(["ID", "Where", "Claim", "Why"],
                    [[r["id"], r.get("location", ""), _clip(r.get("claim", ""), 180),
                      _clip(r.get("explanation", ""), 200)] for r in unv]) if unv else ["None.", ""]

    # --- copy editor ---------------------------------------------------------------------
    L += ["## The copy editor", ""]
    L += [f"{len(entries)} logged change(s). Categories: " + ", ".join(
        f"{CATEGORY_NAMES.get(c.split('(')[0].rstrip('abcdefgh'), c)} ({c}) ×{n}"
        for c, n in Counter(c for e in entries for c in e["categories"]).most_common()) + ".", ""]
    if untagged_fact and not has_fc:
        L += ["### Factual edits (no fact-check report this run)", "",
              "These changes are logged under a factual category (7, or 9g). With no fact-check report they rest "
              "on the copy editor's own knowledge; check each one.", ""]
    elif untagged_fact:
        L += ["### Factual edits NOT tagged [FACT-CHECK]", "",
              "The fact-check report was meant to be the only licence for factual corrections. These changes are "
              "logged under a factual category (7, or 9g) without the tag; check whether each rests on the report "
              "or on the copy editor's memory.", ""]
    if untagged_fact:
        for e in untagged_fact:
            L += [f"- **#{e['n']} [{', '.join(e['categories'])}] {e['location']}:** {e['what']}"
                  + (f" *Why:* {e['why']}" if e["why"] else "")]
        L += [""]
    L += ["### UNVERIFIED notes (doubts left for the author, not edited)", ""]
    L += ([f"- {u}" for u in unverified] or ["None."]) + [""]
    L += ["### Every change, as logged", ""]
    for e in entries:
        tag = "[FACT-CHECK] " if e["fact_check_tagged"] else ""
        L += [f"- **#{e['n']} {tag}[{', '.join(e['categories'])}] {e['location']}:** {e['what']}"
              + (f" *Why:* {e['why']}" if e["why"] else "")]
    L += ["", "The exact before/after text of every change is in the copy-edit diff file beside the guide's "
          "other outputs.", ""]

    # --- the same copy edit without the report --------------------------------------------
    # (S387's one-off comparison; on a run without a fact check it would describe another run)
    cmp_dir = out_dir / "_copy_edit_without_report"
    if has_fc and (cmp_dir / f"psalm_{psalm:03d}_copy_edit_changes.md").exists():
        c_entries, c_unv = parse_change_log(
            (cmp_dir / f"psalm_{psalm:03d}_copy_edit_changes.md").read_text(encoding="utf-8"))
        c_final = (cmp_dir / f"psalm_{psalm:03d}_copy_edited.md").read_text(encoding="utf-8") \
            if (cmp_dir / f"psalm_{psalm:03d}_copy_edited.md").exists() else ""
        c_fact = [e for e in c_entries if is_factual(e)]
        L += ["## The same copy edit without the fact-check report", "",
              "The first attempt's fact check crashed, so the copy editor first ran on the same text with only "
              "the citation report. That edit was kept for comparison; the guide uses the edit made with the "
              "report.", ""]
        L += _table(["", "Without the report", "With the report"], [
            ["Changes logged", len(c_entries), len(entries)],
            ["Factual changes (category 7 or 9g)", len(c_fact), sum(is_factual(e) for e in entries)],
            ["UNVERIFIED notes", len(c_unv), len(unverified)],
            ["Contradicted sentences changed", sum(1 for r in contra if norm(r.get("sentence", ""))
                                                   and norm(r.get("sentence", "")) not in norm(c_final)),
             len(contra) - n_unchanged],
        ])
        if c_fact:
            L += ["Factual changes the copy editor made WITHOUT the report (from its own knowledge):", ""]
            L += [f"- **#{e['n']} [{', '.join(e['categories'])}] {e['location']}:** {e['what']}"
                  + (f" *Why:* {e['why']}" if e["why"] else "") for e in c_fact] + [""]

    # --- free checks -----------------------------------------------------------------------
    L += ["## What was checked for free, by code", ""]
    L += ["### Quoted Bible verses (SQL against the Masoretic text)", ""]
    if cit_pre:
        L += ["Before the copy edit, after the model filter removed false alarms, these quotations did not match "
              "their verse; the list went to the copy editor:", ""]
        L += _table(["Reference", "Where", "Problem"],
                    [[c["ref"], c["location"], _clip(c["issue"], 200)] for c in cit_pre])
    else:
        L += ["Before the copy edit: every checked quotation matched its verse (or the filter judged the "
              "mismatch harmless).", ""]
    if not has_fc:
        L += ["The same check is re-run on the copy-edited text only when the fact check runs, so it did not "
              "run this time.", ""]
        return "\n".join(L).rstrip() + "\n"
    L += ["After the copy edit the same check ran again on the edited text: "
          + (f"{len(cit_post)} new mismatch(es) the copy editor introduced:" if cit_post
             else "the copy editor introduced no new mismatch."), ""]
    if cit_post:
        L += _table(["Reference", "Where", "Problem"],
                    [[c["ref"], c["location"], _clip(c["issue"], 200)] for c in cit_post])
    pc = tele.get("page_check", {})
    L += ["### Web passages checked against their pages", "",
          f"{pc.get('sources', n_sources)} passage(s) gathered; {pc.get('verified', n_verified)} found word for "
          f"word (or nearly) on the page they came from; {pc.get('page_unreadable', 0)} page(s) could not be "
          "fetched or read. A passage not found on its page could not be used to contradict the guide.", ""]
    pre = tele.get("prefetched_evidence", [])
    if pre:
        L += ["### Evidence handed to the fact checker up front", "",
              "Before stage 1 ruled on a chunk, code pulled the commentators' entries on its verses from the "
              "research and the Masoretic text of every verse the chunk cites, so the checker began with "
              "them in hand:", ""]
        L += _table(["Chunk", "Commentator entries", "Verses cited and fetched"],
                    [[p["chunk"], p["commentator_entries"], f"{len(p['cited_verses'])}: "
                      + _clip(", ".join(p["cited_verses"]), 300)] for p in pre])
    tools = Counter(t["tool"] for t in lookups)
    if tools:
        L += ["### Lookups the fact checker chose to make", "",
              "These cost nothing in themselves (the database, the research bundle, Sefaria's free API); the "
              "model pays only for reading the results. By tool: "
              + ", ".join(f"{TOOL_NAMES.get(k, k)} ×{v}" for k, v in tools.most_common())
              + ". Every one is listed in the appendix.", ""]

    # --- supported ---------------------------------------------------------------------------
    sup = [r for r in records if r["verdict"] == "supported"]
    L += ["## Claims the fact checker supported", "",
          f"{len(sup)} claim(s). The evidence column names the source the checker cited.", ""]
    L += _table(["ID", "Where", "Type", "Claim", "Evidence"],
                [[r["id"], r.get("location", ""), r.get("claim_type", ""), _clip(r.get("claim", ""), 200),
                  _clip("; ".join(e.get("source", "") for e in r.get("evidence", [])) or "–", 90)] for r in sup])

    # --- appendix: every lookup ---------------------------------------------------------------
    L += ["## Appendix: every lookup and web action", ""]
    by_run = defaultdict(list)
    for t in trace:
        by_run[t.get("run", "")].append(t)
    for run, items in by_run.items():
        L += [f"### {run} ({items[0].get('model', '')})", ""]
        rows = []
        for t in items:
            if t.get("kind") == "lookup":
                args = ", ".join(f"{v}" for v in (t.get("args") or {}).values())
                rows.append([TOOL_NAMES.get(t.get("tool", ""), t.get("tool", "")), _clip(args, 80),
                             _clip(t.get("found", ""), 260)])
            else:
                q = t.get("query") or ", ".join(t.get("queries") or []) or t.get("url") or t.get("pattern") or ""
                rows.append([f"web {t.get('type', 'search')}", _clip(q, 120), ""])
        L += _table(["Tool", "Asked for", "Found"], rows)

    return "\n".join(L).rstrip() + "\n"


def _cost_section(D: Dict, fc_stage: Optional[Dict]) -> List[str]:
    cost, meta, writer = D["cost"], D["fc"].get("meta", {}), D["writer"]
    stages = cost.get("stages", [])
    total = cost.get("total_cost", 0.0) or 1e-9
    L = ["## What it cost", ""]
    L += ["Every paid step of the run, from the first analysis to the copy edit. Tokens: *fresh input* is "
          "input billed at the full rate; *cached* is input served from a prompt cache at a fraction of it; "
          "*cache write* is input stored in a cache (Anthropic bills it at 1.25× input); *output* includes the "
          "model's reasoning (for OpenAI models the reasoning share is shown separately).", ""]
    last_attempt = max([s.get("attempt", 1) for s in stages] or [1])
    if last_attempt > 1:
        this_run = sum(s["cost_usd"] for s in stages if s.get("attempt", 1) == last_attempt)
        L += [f"This psalm's cost file holds {last_attempt} runs (a resumed or partial re-run continues it), and "
              f"every stage is marked with the attempt it belongs to. **This run (attempt {last_attempt}) cost "
              f"{_usd(this_run)}**; the total below is all runs together. A stage this run reused from an earlier "
              "one cost nothing and has no row of its own.", ""]
    rows = []
    for s in stages:
        ms = s["models"]
        if not ms and not s["charges"]:
            continue
        agg = Counter()
        for m in ms.values():
            for k in ("call_count", "input_tokens", "cache_read_tokens", "cache_write_tokens",
                      "cache_write_1h_tokens", "output_tokens", "thinking_tokens"):
                agg[k] += m.get(k, 0)
        out = agg["output_tokens"] + agg["thinking_tokens"]
        rsn = f" ({agg['thinking_tokens']:,} reasoning)" if agg["thinking_tokens"] else ""
        fee = sum(c["usd"] for c in s["charges"])
        name = s["stage"] + (f" (attempt {s.get('attempt', 1)})" if last_attempt > 1 else "")
        rows.append([name, ", ".join(ms) or "–", agg["call_count"], _k(agg["input_tokens"]),
                     _k(agg["cache_read_tokens"]), _k(agg["cache_write_tokens"] + agg["cache_write_1h_tokens"]),
                     _k(out) + rsn, (_usd(fee) if fee else "–"), _usd(s["cost_usd"]),
                     f"{100 * s['cost_usd'] / total:.0f}%"])
    staged = sum(s["cost_usd"] for s in stages)
    if total - staged > 0.005:
        rows.append(["(not inside a recorded stage)", "", "", "", "", "", "", "", _usd(total - staged),
                     f"{100 * (total - staged) / total:.0f}%"])
    rows.append(["**Total**", "", "", "", "", "", "", "", f"**{_usd(total)}**", "100%"])
    L += _table(["Stage", "Model(s)", "Calls", "Fresh input", "Cached", "Cache write", "Output",
                 "Non-token fees", "Cost", "Share"], rows)
    for note in cost.get("notes", []):
        L += [f"- {note}"]
    if cost.get("notes"):
        L += [""]

    if writer.get("calls"):
        L += ["### Inside the writer", "",
              f"{writer.get('model', '')} at reasoning effort {writer.get('effort', '')}. Call 1 writes the essay "
              "and stores the dossier in a 5-minute cache; call 2 (liturgy and verse commentary) reads it back.", ""]
        L += _table(["Call", "Seconds", "Fresh input", "Cache read", "Cache write", "Output (incl. reasoning)", "Cost"],
                    [[c["label"], c["seconds"], _k(c["usage"]["input"]), _k(c["usage"]["cache_read"]),
                      _k(c["usage"]["cache_write"] + c["usage"].get("cache_write_1h", 0)), _k(c["usage"]["output"]),
                      _usd(c["cost_usd"])] for c in writer["calls"]])
        if writer.get("structure_problems") or writer.get("essay_problems"):
            L += ["Structure problems the writer's output check reported: "
                  + "; ".join(writer.get("essay_problems", []) + writer.get("structure_problems", [])), ""]

    ps = meta.get("stages", {}).get("per_stage", {})
    if ps:
        L += ["### Inside the fact check", ""]
        names = {"local": "stage 1: local evidence, no web", "gather": "stage 2a: gather web passages",
                 "judge": "stage 2b: judge from checked passages", "review": "stage 3: review",
                 "web": "stage 2: web"}
        L += _table(["Step", "Model", "Fresh input", "Cached", "Output", "of which reasoning", "Searches",
                     "Lookups", "Cost"],
                    # "Fresh" = first-sight input: plain + written to cache (1.25x on GPT-5.6+, S388)
                    [[names.get(k, k), v["model"], _k(v["usage"]["input"] + v["usage"].get("cache_write", 0)),
                      _k(v["usage"]["cached"]),
                      _k(v["usage"]["output"] + v["usage"]["reasoning"]), _k(v["usage"]["reasoning"]),
                      v["searches"], sum(v.get("tools", {}).values()), _usd(v["cost_usd"])]
                     for k, v in ps.items()])
        for fch in D["tele"].get("failed_chunks", []):
            L += [f"- A chunk FAILED and its claims were not checked: {fch['chunk']} ({fch['error']})"]
        trunc = [r for r in D["tele"].get("runs", []) if r.get("truncated")]
        if trunc:
            L += [f"- {len(trunc)} call(s) returned a truncated answer ("
                  + ", ".join(r["label"] for r in trunc) + "); their chunks were re-checked in halves."]
        L += [f"Web-search fees: {meta.get('web_searches', 0)} searches × $0.01 = "
              f"{_usd(meta.get('search_cost_usd', 0.0))}, included above (search results are billed as input "
              f"tokens as well). Fact check total {_usd(meta.get('cost_usd', 0.0))} in "
              f"{round(meta.get('seconds', 0) / 60)} minutes.", ""]

    # A promo flag saved at run time is only as good as the price table was then. Session 388:
    # gpt-6-sol's "promo" turned out to be its standard price, so a flag is shown only while the
    # model is STILL on introductory pricing in today's table.
    from src.utils.cost_tracker import INTRO_PRICING
    promo = [(s["stage"], m, r) for s in stages for m, r in s["models"].items()
             if r.get("promo_through") and m in INTRO_PRICING]
    if promo:
        now = sum(r["cost_usd"] for _, _, r in promo)
        dur = sum(r.get("cost_usd_at_durable_rates", r["cost_usd"]) for _, _, r in promo)
        L += ["### Introductory prices", "",
              "Some of these models are on a launch price that ends on a known date. At the durable price the "
              "same tokens would cost:", ""]
        L += _table(["Stage", "Model", "Promo until", "Cost now", "At durable rates"],
                    [[st, m, r["promo_through"], _usd(r["cost_usd"]), _usd(r.get("cost_usd_at_durable_rates", 0))]
                     for st, m, r in promo])
        L += [f"Run total at durable rates: {_usd(total - now + dur)} (now {_usd(total)}).", ""]
    return L


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def render_docx(markdown: str, out_path: Path, psalm: int) -> Path:
    """Render the report through the guide's own Hebrew-aware paragraph machinery, plus tables."""
    from docx import Document
    from docx.shared import Inches, Pt

    from src.utils.divine_names_modifier import DivineNamesModifier
    from src.utils.document_generator import DocumentGenerator, add_page_number

    g = DocumentGenerator.__new__(DocumentGenerator)
    g.document = Document()
    g.modifier = DivineNamesModifier()
    g.psalm_num = psalm
    g._set_default_styles()
    styles = g.document.styles
    styles['BodySans'].font.size = Pt(10.5)
    g._set_style_complex_size(styles['BodySans'], 11.5)
    for s in g.document.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.75)

    def cell_text(cell, text, bold=False):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        g._process_markdown_formatting(p, f"**{text}**" if bold and text else text)
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.name = 'Aptos'

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split(" | ")] for r in block
                    if not re.match(r"^\|(-{3}\|)+\s*$", r.strip())]
            ncol = max(len(r) for r in rows)
            t = g.document.add_table(rows=len(rows), cols=ncol)
            t.style = 'Table Grid'
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell_text(t.cell(ri, ci), r[ci] if ci < len(r) else "", bold=(ri == 0))
            g.document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if line.startswith("# "):
            g.document.add_heading(line[2:].strip(), level=1)
        elif not line.strip():
            pass
        else:
            g._add_paragraph_with_markdown(line, style='BodySans')
        i += 1
    from docx.oxml.ns import qn
    for tag, val in (("w:sz", 21), ("w:szCs", 23)):   # list items carry an explicit 12pt
        for el in g.document.element.body.iter(qn(tag)):
            if int(el.get(qn("w:val"))) > val:
                el.set(qn("w:val"), str(val))
    p = g.document.sections[0].footer.paragraphs[0]
    add_page_number(p)
    g._join_rtl_runs_across_whitespace()
    g._mirror_bold_to_complex_script()
    g._fix_complex_script_fonts()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    g.document.save(out_path)
    return out_path


def build_editors_report(psalm: int, out_dir: Path, docx_path: Optional[Path] = None) -> Path:
    out_dir = Path(out_dir)
    md = build_markdown(psalm, out_dir)
    md_path = out_dir / f"psalm_{psalm:03d}_editors_report.md"
    md_path.write_text(md, encoding="utf-8")
    return render_docx(md, docx_path or out_dir / f"psalm_{psalm:03d}_editors_report.docx", psalm)


if __name__ == "__main__":
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    import argparse
    ap = argparse.ArgumentParser(description="Build the editors' report for a finished psalm ($0).")
    ap.add_argument("psalm", type=int)
    ap.add_argument("--output-dir", default=None)
    ap.add_argument("--docx", default=None, help="where to write the DOCX (default: beside the outputs)")
    a = ap.parse_args()
    from src.utils.debug_paths import psalm_output_dir
    od = Path(a.output_dir) if a.output_dir else psalm_output_dir(a.psalm)
    print(build_editors_report(a.psalm, od, Path(a.docx) if a.docx else None))
